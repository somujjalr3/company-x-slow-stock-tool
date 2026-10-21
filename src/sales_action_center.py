"""
Sales Action Center Module for Company X Slow Stock Action Tool.

Answers the primary operational question for commercial sales representatives:
"Who should the salesperson call today and what should they sell?"

Key Capabilities:
1. generate_customer_sku_action():
   Produces deterministic, explainable action recommendations for each Customer-SKU combination.
   Example:
     "Contact Customer 220 and promote SKU1005 because the customer previously purchased the item and has ₹665,004 of 6M purchase value."

2. build_sales_action_table():
   Consolidates actionable Push SKUs and their pre-qualified customer targets
   into a single, unified cross-portfolio action ledger.

3. filter_sales_action_table():
   Multi-parameter filtering across:
   - SKU
   - Category
   - Priority
   - Risk Level
   - Customer Bucket
   - Industry
   - Customer ID

4. sort_sales_action_table():
   Deterministic multi-level sorting with Priority Score as primary sort key.

5. export_sales_action_excel():
   Professional, executive-styled Excel export of the sales action call list.

CONSTRAINTS & INVARIANTS:
- Does NOT alter Rules R1-R7 or customer bucket logic.
- 100% deterministic, explainable calculations without AI.
"""

import io
from typing import Dict, List, Tuple, Any, Optional
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from src.logger import get_logger
from src.sales_priority import calculate_sales_priority

logger = get_logger("SalesActionCenter")


def generate_customer_sku_action(
    customer_name: str,
    sku: str,
    bucket: int,
    spend: float,
    industry: str = None,
    category: str = None,
    bought_switch: bool = False,
    switch_skus: list = None,
) -> str:
    """
    Generate deterministic clear action recommendation for an individual customer-SKU combination.

    Example:
      'Contact Customer 220 and promote SKU1005 because the customer previously purchased the item and has ₹665,004 of 6M purchase value.'
    """
    spend_val = float(spend) if pd.notna(spend) else 0.0
    spend_str = f"₹{spend_val:,.0f}"
    cust_str = str(customer_name).strip() if (customer_name and pd.notna(customer_name) and str(customer_name).strip() != "") else "the customer"
    sku_str = str(sku).strip().upper()

    try:
        b = int(bucket) if pd.notna(bucket) else 1
    except (ValueError, TypeError):
        b = 1

    if b == 1:
        if bought_switch and switch_skus:
            sw_text = f" (switch partner {', '.join(switch_skus)})"
            return f"Contact {cust_str} and promote {sku_str} because the customer previously purchased the switch partner item{sw_text} and has {spend_str} of 6M purchase value."
        else:
            return f"Contact {cust_str} and promote {sku_str} because the customer previously purchased the item and has {spend_str} of 6M purchase value."
    elif b == 2:
        ind_label = str(industry).strip() if (industry and pd.notna(industry) and str(industry).strip() != "") else "a peer"
        return f"Contact {cust_str} and promote {sku_str} because the customer operates in the {ind_label} industry associated with previous buyers and has {spend_str} of 6M purchase value."
    elif b == 3:
        cat_label = str(category).strip() if (category and pd.notna(category) and str(category).strip() != "") else "the same"
        return f"Contact {cust_str} and promote {sku_str} because the customer previously purchased other items in the {cat_label} category and has {spend_str} of 6M purchase value."
    else:
        return f"Contact {cust_str} and promote {sku_str} because the customer is an eligible account with {spend_str} of 6M purchase value."


def build_sales_action_table(
    items_df: pd.DataFrame,
    targets_dict: dict,
    sales_trends: dict = None,
) -> pd.DataFrame:
    """
    Constructs the consolidated cross-portfolio Sales Action Table.

    Combines actionable Push SKUs with their qualified customer targets across Buckets 1, 2, and 3.

    Required columns produced:
      - Priority (SKU Priority Level, e.g. 'Critical', 'High', 'Medium', 'Low')
      - SKU (Catalog SKU)
      - Item Description (Item Name)
      - Customer ID (CustomerID code, e.g. 'C220')
      - Customer Industry (Industry classification)
      - Bucket (R6 Bucket number: 1, 2, 3)
      - 6M Purchase Value (Total 6-month spend in ₹)
      - Opportunity Score (0-100 deterministic account score)
      - Aged Value (Capital trapped in 91+ day aged inventory)
      - Months to Clear (Runway in months)
      - Recommended Action (Operational call/promote recommendation)
      - Reason (Underlying customer qualification reason)

    Supporting metadata columns:
      - Priority Score (0-100 numeric score for sorting)
      - Priority Level ('Critical' | 'High' | 'Medium' | 'Low')
      - Category (Item Category)
      - Risk Level (Inventory Risk Tier)
      - Risk Score (Inventory Risk Score)
      - Customer Name (Full customer name)
      - Customer Bucket (Display string, e.g. 'Bucket 1')
      - Bucket Number (Integer 1, 2, 3)
      - Opportunity Level ('Prime' | 'High' | 'Solid' | 'Developing')
      - Relevant Purchase Evidence (Transaction details)
    """
    if items_df is None or items_df.empty or not targets_dict:
        return pd.DataFrame(columns=[
            "Priority", "SKU", "Item Description", "Customer ID", "Customer Industry",
            "Bucket", "6M Purchase Value", "Opportunity Score", "Aged Value",
            "Months to Clear", "Recommended Action", "Reason"
        ])

    items = items_df.copy()
    if "Priority Score" not in items.columns or "Priority Level" not in items.columns:
        items = calculate_sales_priority(items, sales_trends=sales_trends)

    sku_meta = items.set_index("SKU")

    rows = []
    # Only iterate through Push items
    push_skus = [s for s in targets_dict.keys() if s in sku_meta.index and sku_meta.loc[s, "Class"] == "Push"]

    for sku in push_skus:
        t_df = targets_dict.get(sku)
        if t_df is None or t_df.empty:
            continue

        item_row = sku_meta.loc[sku]
        item_name = str(item_row.get("Item Name", sku))
        category = str(item_row.get("Category", "Unassigned"))
        aged_val = float(item_row.get("Aged Value", 0.0))
        mtc = float(item_row.get("Months to Clear", 0.0))
        risk_score = int(item_row.get("Risk Score", 0))
        risk_lvl = str(item_row.get("Risk Level", "Medium"))
        p_score = int(item_row.get("Priority Score", 0))
        p_lvl = str(item_row.get("Priority Level", "Medium"))
        why_priority = str(item_row.get("Why Priority", ""))

        for _, cust in t_df.iterrows():
            cid = str(cust.get("CustomerID") or cust.get("Customer ID", "")).strip()
            cname = str(cust.get("CustomerName") or cust.get("Customer Name", cid)).strip()
            industry = str(cust.get("Industry", "")).strip() if pd.notna(cust.get("Industry")) else "Unassigned"
            if industry in ["", "nan", "None", "NA", "<NA>"]:
                industry = "Unassigned"

            b_raw = cust.get("Bucket", 1)
            try:
                b_num = int(b_raw)
            except (ValueError, TypeError):
                b_num = 1

            spend = float(cust.get("6M Purchase Value", 0.0))
            opp_score = int(cust.get("Opportunity Score", 0))
            opp_lvl = str(cust.get("Opportunity Level", "Developing"))
            evidence = str(cust.get("Relevant Purchase Evidence", ""))
            reason = str(cust.get("Reason") or cust.get("Why this customer?", ""))

            # Generate specific action recommendation
            action_rec = generate_customer_sku_action(
                customer_name=cname,
                sku=sku,
                bucket=b_num,
                spend=spend,
                industry=industry,
                category=category,
                bought_switch="switch partner" in evidence.lower(),
            )

            rows.append({
                # Core 12 specified columns
                "Priority": p_lvl,
                "SKU": sku,
                "Item Description": item_name,
                "Customer ID": cid,
                "Customer Industry": industry,
                "Bucket": b_num,
                "6M Purchase Value": spend,
                "Opportunity Score": opp_score,
                "Aged Value": aged_val,
                "Months to Clear": mtc,
                "Recommended Action": action_rec,
                "Reason": reason,

                # Supporting / Sorting / Filtering columns
                "Priority Score": p_score,
                "Priority Level": p_lvl,
                "Why Priority": why_priority,
                "Category": category,
                "Risk Level": risk_lvl,
                "Risk Score": risk_score,
                "Customer Name": cname,
                "CustomerName": cname,
                "CustomerID": cid,
                "Industry": industry,
                "Customer Bucket": f"Bucket {b_num}",
                "Bucket Number": b_num,
                "Bucket Name": str(cust.get("Bucket Name", f"Bucket {b_num}")),
                "Opportunity Level": opp_lvl,
                "Relevant Purchase Evidence": evidence,
                "R7 Priority": int(cust.get("Priority", 1)),
            })

    df_out = pd.DataFrame(rows)
    if df_out.empty:
        return df_out

    # Default sort: Priority Score descending, Bucket ascending, 6M Spend descending
    df_out = df_out.sort_values(
        by=["Priority Score", "Bucket Number", "6M Purchase Value"],
        ascending=[False, True, False]
    ).reset_index(drop=True)

    logger.info(
        "build_sales_action_table complete: %d total sales action recommendations across %d Push SKUs",
        len(df_out), len(df_out["SKU"].unique())
    )
    return df_out


def filter_sales_action_table(
    df: pd.DataFrame,
    sku=None,
    category=None,
    priority=None,
    risk_level=None,
    customer_bucket=None,
    industry=None,
    customer_id=None,
) -> pd.DataFrame:
    """
    Applies multi-parameter filtering to the Sales Action Table.

    All filters default to None ('All') and combine via logical AND.
    """
    if df is None or df.empty:
        return pd.DataFrame()

    filtered = df.copy()

    # 1. SKU filter
    if sku is not None and str(sku).strip().lower() not in ["all", "all skus", ""]:
        if isinstance(sku, (list, tuple, set)):
            s_set = {str(s).strip().upper() for s in sku if str(s).strip() and str(s).strip().lower() not in ["all", "all skus"]}
            if s_set:
                filtered = filtered[filtered["SKU"].str.upper().isin(s_set)]
        else:
            s_val = str(sku).strip().upper()
            filtered = filtered[filtered["SKU"].str.upper() == s_val]

    # 2. Category filter
    if category is not None and str(category).strip().lower() not in ["all", "all categories", ""]:
        if isinstance(category, (list, tuple, set)):
            c_set = {str(c).strip().lower() for c in category if str(c).strip() and str(c).strip().lower() not in ["all", "all categories"]}
            if c_set:
                filtered = filtered[filtered["Category"].str.lower().isin(c_set)]
        else:
            c_val = str(category).strip().lower()
            filtered = filtered[filtered["Category"].str.lower() == c_val]

    # 3. Priority filter
    if priority is not None and str(priority).strip().lower() not in ["all", "all priorities", ""]:
        if isinstance(priority, (list, tuple, set)):
            p_set = {str(p).strip().lower() for p in priority if str(p).strip() and str(p).strip().lower() not in ["all", "all priorities"]}
            if p_set:
                filtered = filtered[
                    filtered["Priority Level"].str.lower().isin(p_set) |
                    filtered["Priority"].astype(str).str.lower().isin(p_set)
                ]
        else:
            p_val = str(priority).strip().lower()
            filtered = filtered[
                (filtered["Priority Level"].str.lower() == p_val) |
                (filtered["Priority"].astype(str).str.lower() == p_val)
            ]

    # 4. Risk Level filter
    if risk_level is not None and str(risk_level).strip().lower() not in ["all", "all risk levels", ""]:
        if isinstance(risk_level, (list, tuple, set)):
            r_set = {str(r).strip().lower() for r in risk_level if str(r).strip() and str(r).strip().lower() not in ["all", "all risk levels"]}
            if r_set:
                filtered = filtered[filtered["Risk Level"].str.lower().isin(r_set)]
        else:
            r_val = str(risk_level).strip().lower()
            filtered = filtered[filtered["Risk Level"].str.lower() == r_val]

    # 5. Customer Bucket filter
    if customer_bucket is not None and str(customer_bucket).strip().lower() not in ["all", "all buckets", ""]:
        def _parse_b(b):
            s = str(b).strip()
            digits = "".join(ch for ch in s if ch.isdigit())
            return int(digits) if digits else None

        if isinstance(customer_bucket, (list, tuple, set)):
            target_buckets = {_parse_b(b) for b in customer_bucket}
            target_buckets.discard(None)
            if target_buckets:
                filtered = filtered[
                    filtered["Bucket"].isin(target_buckets) |
                    filtered["Bucket Number"].isin(target_buckets)
                ]
        else:
            tb = _parse_b(customer_bucket)
            if tb is not None:
                filtered = filtered[(filtered["Bucket"] == tb) | (filtered["Bucket Number"] == tb)]

    # 6. Industry filter
    if industry is not None and str(industry).strip().lower() not in ["all", "all industries", ""]:
        if isinstance(industry, (list, tuple, set)):
            ind_set = {str(i).strip().lower() for i in industry if str(i).strip() and str(i).strip().lower() not in ["all", "all industries"]}
            if ind_set:
                filtered = filtered[filtered["Customer Industry"].astype(str).str.lower().isin(ind_set)]
        else:
            i_val = str(industry).strip().lower()
            filtered = filtered[filtered["Customer Industry"].astype(str).str.lower() == i_val]

    # 7. Customer ID filter
    if customer_id is not None and str(customer_id).strip().lower() not in ["all", "all customers", ""]:
        if isinstance(customer_id, (list, tuple, set)):
            cid_set = {str(c).strip().upper() for c in customer_id if str(c).strip() and str(c).strip().lower() not in ["all", "all customers"]}
            if cid_set:
                filtered = filtered[
                    filtered["Customer ID"].str.upper().isin(cid_set) |
                    filtered["CustomerID"].str.upper().isin(cid_set) |
                    filtered["Customer Name"].str.upper().isin(cid_set)
                ]
        else:
            cid_val = str(customer_id).strip().upper()
            filtered = filtered[
                (filtered["Customer ID"].str.upper() == cid_val) |
                (filtered["CustomerID"].str.upper() == cid_val) |
                (filtered["Customer Name"].str.upper() == cid_val) |
                (filtered["Customer ID"].str.upper().str.contains(cid_val, na=False)) |
                (filtered["Customer Name"].str.upper().str.contains(cid_val, na=False))
            ]

    return filtered.reset_index(drop=True)


def sort_sales_action_table(
    df: pd.DataFrame,
    sort_by: str = "Priority Score",
    ascending: bool = False,
) -> pd.DataFrame:
    """
    Sorts the Sales Action Table with deterministic multi-level criteria.
    Defaults to Priority Score descending.
    """
    if df is None or df.empty:
        return pd.DataFrame()

    sorted_df = df.copy()
    s = str(sort_by).strip().lower()

    if "priority score" in s:
        sorted_df = sorted_df.sort_values(
            by=["Priority Score", "Bucket Number", "6M Purchase Value"],
            ascending=[ascending, True, False]
        )
    elif "opportunity score" in s:
        sorted_df = sorted_df.sort_values(
            by=["Opportunity Score", "Priority Score", "6M Purchase Value"],
            ascending=[ascending, False, False]
        )
    elif "aged value" in s:
        sorted_df = sorted_df.sort_values(
            by=["Aged Value", "Priority Score", "6M Purchase Value"],
            ascending=[ascending, False, False]
        )
    elif "spend" in s or "purchase value" in s:
        sorted_df = sorted_df.sort_values(
            by=["6M Purchase Value", "Priority Score"],
            ascending=[ascending, False]
        )
    elif "bucket" in s:
        b_asc = True if ascending is None else ascending
        sorted_df = sorted_df.sort_values(
            by=["Bucket Number", "Priority Score", "6M Purchase Value"],
            ascending=[b_asc, False, False]
        )
    elif "months" in s:
        sorted_df = sorted_df.sort_values(
            by=["Months to Clear", "Priority Score"],
            ascending=[ascending, False]
        )
    else:
        sorted_df = sorted_df.sort_values(
            by=["Priority Score", "Bucket Number", "6M Purchase Value"],
            ascending=[False, True, False]
        )

    return sorted_df.reset_index(drop=True)


def export_sales_action_excel(df: pd.DataFrame) -> bytes:
    """
    Generates a professionally formatted Excel workbook (.xlsx) for the Sales Action List.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales Action List"

    # Styling Constants
    font_title = Font(name="Calibri", size=14, bold=True, color="0F294A")
    font_sub = Font(name="Calibri", size=9, italic=True, color="475569")
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=9, color="000000")
    header_fill = PatternFill(start_color="0F294A", end_color="0F294A", fill_type="solid")
    border_light = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    # Title Banner
    ws["A1"] = "Company X — Sales Action Center Call List"
    ws["A1"].font = font_title
    ws["A2"] = "Actionable Slow-Stock Customer Targets | Strictly Deterministic Rules R1–R7"
    ws["A2"].font = font_sub

    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 16
    ws.row_dimensions[3].height = 10

    export_cols = [
        ("Priority", "Priority"),
        ("Priority Score", "Priority Score"),
        ("SKU", "SKU"),
        ("Item Description", "Item Description"),
        ("Category", "Category"),
        ("Customer ID", "Customer ID"),
        ("Customer Name", "Customer Name"),
        ("Customer Industry", "Customer Industry"),
        ("Bucket", "Bucket"),
        ("6M Purchase Value", "6M Purchase Value"),
        ("Opportunity Score", "Opportunity Score"),
        ("Opportunity Level", "Opportunity Level"),
        ("Aged Value", "Aged Value"),
        ("Months to Clear", "Months to Clear"),
        ("Recommended Action", "Recommended Action"),
        ("Reason", "Reason"),
    ]

    header_row = 4
    ws.row_dimensions[header_row].height = 24

    for col_idx, (col_title, _) in enumerate(export_cols, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=col_title)
        cell.font = font_header
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_light

    if df is not None and not df.empty:
        for r_idx, (_, row) in enumerate(df.iterrows(), start=header_row + 1):
            ws.row_dimensions[r_idx].height = 20
            for c_idx, (_, df_col) in enumerate(export_cols, start=1):
                val = row.get(df_col, "")
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.font = font_data
                cell.border = border_light

                # Format specific data types
                if df_col in ["6M Purchase Value", "Aged Value"]:
                    cell.value = float(val) if pd.notna(val) else 0.0
                    cell.number_format = "₹#,##0"
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif df_col == "Months to Clear":
                    if pd.isna(val) or val == float("inf") or np.isinf(val):
                        cell.value = "∞"
                    else:
                        cell.value = round(float(val), 2)
                        cell.number_format = "0.00"
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif df_col in ["Priority Score", "Opportunity Score", "Bucket"]:
                    cell.value = int(val) if pd.notna(val) else 0
                    cell.number_format = "0"
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif df_col in ["SKU", "Customer ID", "Priority", "Opportunity Level"]:
                    cell.value = str(val) if pd.notna(val) else ""
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.value = str(val) if pd.notna(val) else ""
                    cell.alignment = Alignment(horizontal="left", vertical="center")

    # Auto-fit columns
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            if cell.row < header_row:
                continue
            v = str(cell.value or "")
            if len(v) > max_len:
                max_len = len(v)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 11), 50)

    # Freeze panes below header
    ws.freeze_panes = "A5"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ==============================================================================
# SALES OUTREACH & EXECUTION TRACKER (Item 26)
# ==============================================================================
TRACKER_STATUS_OPTIONS = [
    "Not Contacted",
    "Contacted / In Discussion",
    "Follow Up Scheduled",
    "Order Placed (Liquidated)",
    "Customer Declined",
]

def init_sales_tracker(sales_action_df: pd.DataFrame, existing_records: Optional[Dict[str, Dict[str, Any]]] = None) -> pd.DataFrame:
    """
    Initializes or refreshes the operational Sales Outreach & Execution Tracker.
    Preserves any existing manual call logs and dispositions across page reruns.
    """
    if sales_action_df is None or sales_action_df.empty:
        return pd.DataFrame(columns=[
            "Lead ID", "Priority", "Priority Score", "SKU", "Item Description",
            "Customer ID", "Customer Name", "Customer Industry", "Bucket",
            "Opportunity Score", "6M Purchase Value", "Aged Value", "Months to Clear",
            "Recommended Action", "Call Status", "Last Contact Date",
            "Units Committed", "Value Realized (Rs)", "Sales Rep Notes"
        ])

    df = sales_action_df.copy()
    existing_map = existing_records or {}

    rows = []
    for _, r in df.iterrows():
        sku = str(r.get("SKU", "")).strip()
        cid = str(r.get("Customer ID", "")).strip()
        lead_id = f"{sku}::{cid}"

        # Check existing state
        stored = existing_map.get(lead_id, {})
        status = stored.get("Call Status", "Not Contacted")
        contact_date = stored.get("Last Contact Date", "-")
        units = int(stored.get("Units Committed", 0))
        val = float(stored.get("Value Realized (Rs)", 0.0))
        notes = str(stored.get("Sales Rep Notes", ""))

        rows.append({
            "Lead ID": lead_id,
            "Priority": str(r.get("Priority", "Medium")),
            "Priority Score": int(r.get("Priority Score", 50)),
            "SKU": sku,
            "Item Description": str(r.get("Item Description", "")),
            "Customer ID": cid,
            "Customer Name": str(r.get("Customer Name", "")),
            "Customer Industry": str(r.get("Customer Industry", "Unassigned")),
            "Bucket": int(r.get("Bucket Number", r.get("Bucket", 1))),
            "Opportunity Score": int(r.get("Opportunity Score", 50)),
            "6M Purchase Value": float(r.get("6M Purchase Value", 0.0)),
            "Aged Value": float(r.get("Aged Value", 0.0)),
            "Months to Clear": float(r.get("Months to Clear", 0.0)) if pd.notna(r.get("Months to Clear")) else float("inf"),
            "Recommended Action": str(r.get("Recommended Action", "")),
            "Call Status": status,
            "Last Contact Date": contact_date,
            "Units Committed": units,
            "Value Realized (Rs)": val,
            "Sales Rep Notes": notes,
        })

    return pd.DataFrame(rows)


def update_sales_tracker_lead(
    tracker_df: pd.DataFrame,
    sku: str,
    customer_id: str,
    status: str,
    notes: str = "",
    units: int = 0,
    value: float = 0.0,
    contact_date: Optional[str] = None
) -> pd.DataFrame:
    """
    Updates the outreach disposition and outcome for a specific SKU-Customer lead.
    """
    if tracker_df is None or tracker_df.empty:
        return tracker_df

    df = tracker_df.copy()
    lead_id = f"{str(sku).strip()}::{str(customer_id).strip()}"
    idx = df[df["Lead ID"] == lead_id].index

    if len(idx) > 0:
        target_idx = idx[0]
        df.at[target_idx, "Call Status"] = status
        df.at[target_idx, "Sales Rep Notes"] = str(notes).strip()
        df.at[target_idx, "Units Committed"] = int(units)
        df.at[target_idx, "Value Realized (Rs)"] = float(value)
        if contact_date:
            df.at[target_idx, "Last Contact Date"] = str(contact_date)
        elif df.at[target_idx, "Last Contact Date"] == "-":
            df.at[target_idx, "Last Contact Date"] = pd.Timestamp.now().strftime("%Y-%m-%d")
        logger.info(f"Updated sales tracker for lead {lead_id} -> {status}")
    else:
        logger.warning(f"Lead ID {lead_id} not found in tracker table.")

    return df


def calculate_sales_tracker_kpis(tracker_df: pd.DataFrame) -> Dict[str, Any]:
    """
    Computes real-time executive and operational outreach KPIs.
    """
    if tracker_df is None or tracker_df.empty:
        return {
            "total_leads": 0,
            "contacted_count": 0,
            "pending_count": 0,
            "in_discussion_count": 0,
            "orders_placed_count": 0,
            "declined_count": 0,
            "conversion_rate_pct": 0.0,
            "total_aged_target": 0.0,
            "liquidated_value": 0.0,
            "liquidated_units": 0,
            "active_skus_tracked": 0,
            "active_accounts_tracked": 0,
        }

    total = len(tracker_df)
    contacted = int((tracker_df["Call Status"] != "Not Contacted").sum())
    pending = int((tracker_df["Call Status"] == "Not Contacted").sum())
    in_disc = int(tracker_df["Call Status"].str.contains("In Discussion|Follow Up", na=False).sum())
    orders = int((tracker_df["Call Status"] == "Order Placed (Liquidated)").sum())
    declined = int((tracker_df["Call Status"] == "Customer Declined").sum())

    conversion_rate = (orders / contacted * 100.0) if contacted > 0 else 0.0
    tot_aged = float(tracker_df.drop_duplicates("SKU")["Aged Value"].sum())
    liq_val = float(tracker_df["Value Realized (Rs)"].sum())
    liq_units = int(tracker_df["Units Committed"].sum())

    return {
        "total_leads": total,
        "contacted_count": contacted,
        "pending_count": pending,
        "in_discussion_count": in_disc,
        "orders_placed_count": orders,
        "declined_count": declined,
        "conversion_rate_pct": conversion_rate,
        "total_aged_target": tot_aged,
        "liquidated_value": liq_val,
        "liquidated_units": liq_units,
        "active_skus_tracked": int(tracker_df["SKU"].nunique()),
        "active_accounts_tracked": int(tracker_df["Customer ID"].nunique()),
    }


def export_sales_tracker_excel(tracker_df: pd.DataFrame) -> bytes:
    """
    Exports the Sales Execution Tracker ledger to an executive Excel spreadsheet.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales Execution Tracker"

    # Styling
    font_title = Font(name="Calibri", size=14, bold=True, color="0F294A")
    font_sub = Font(name="Calibri", size=9, italic=True, color="475569")
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=9, color="000000")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    border_light = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    ws["A1"] = "Company X — Sales Outreach & Execution Tracker"
    ws["A1"].font = font_title
    ws["A2"] = "Commercial Execution Log | Daily Call Dispositions & Slow-Stock Liquidation"
    ws["A2"].font = font_sub

    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 16
    ws.row_dimensions[3].height = 10

    cols = [
        ("Lead ID", "Lead ID"),
        ("Call Status", "Call Status"),
        ("Last Contact Date", "Last Contact Date"),
        ("Units Committed", "Units Committed"),
        ("Value Realized (Rs)", "Value Realized (Rs)"),
        ("Sales Rep Notes", "Sales Rep Notes"),
        ("Priority", "Priority"),
        ("Priority Score", "Priority Score"),
        ("SKU", "SKU"),
        ("Item Description", "Item Description"),
        ("Customer ID", "Customer ID"),
        ("Customer Name", "Customer Name"),
        ("Customer Industry", "Customer Industry"),
        ("Bucket", "Bucket"),
        ("Opportunity Score", "Opportunity Score"),
        ("6M Purchase Value", "6M Purchase Value"),
        ("Aged Value", "Aged Value"),
        ("Months to Clear", "Months to Clear"),
        ("Recommended Action", "Recommended Action"),
    ]

    header_row = 4
    ws.row_dimensions[header_row].height = 24

    for c_idx, (col_title, _) in enumerate(cols, start=1):
        cell = ws.cell(row=header_row, column=c_idx, value=col_title)
        cell.font = font_header
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_light

    if tracker_df is not None and not tracker_df.empty:
        for r_idx, (_, row) in enumerate(tracker_df.iterrows(), start=header_row + 1):
            ws.row_dimensions[r_idx].height = 20
            for c_idx, (_, df_col) in enumerate(cols, start=1):
                val = row.get(df_col, "")
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.font = font_data
                cell.border = border_light

                if df_col in ["Value Realized (Rs)", "6M Purchase Value", "Aged Value"]:
                    cell.value = float(val) if pd.notna(val) else 0.0
                    cell.number_format = "₹#,##0"
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif df_col == "Months to Clear":
                    if pd.isna(val) or val == float("inf") or np.isinf(val):
                        cell.value = "∞"
                    else:
                        cell.value = round(float(val), 2)
                        cell.number_format = "0.00"
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif df_col in ["Units Committed", "Priority Score", "Opportunity Score", "Bucket"]:
                    cell.value = int(val) if pd.notna(val) else 0
                    cell.number_format = "#,##0" if df_col == "Units Committed" else "0"
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif df_col == "Call Status":
                    cell.value = str(val)
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    # Color status
                    st_str = str(val)
                    if "Order Placed" in st_str:
                        cell.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
                        cell.font = Font(name="Calibri", size=9, bold=True, color="166534")
                    elif "Contacted" in st_str or "Discussion" in st_str:
                        cell.fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
                        cell.font = Font(name="Calibri", size=9, bold=True, color="1E40AF")
                    elif "Declined" in st_str:
                        cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
                        cell.font = Font(name="Calibri", size=9, bold=True, color="991B1B")
                else:
                    cell.value = str(val) if pd.notna(val) else ""
                    cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            if cell.row < header_row:
                continue
            v = str(cell.value or "")
            if len(v) > max_len:
                max_len = len(v)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    ws.freeze_panes = "A5"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
