"""
Unit tests for Dead Stock Analysis module (src/dead_stock.py).

Verifies:
1. Dead Stock identification condition: 6M Sales Qty = 0.
2. All 10 required columns are present:
   - SKU
   - Description
   - Category
   - Total Stock
   - Inventory Value
   - Aged Qty
   - Aged Value
   - Last Available Sales Information
   - Risk Score
   - Recommended Action
3. Clear separation of Dead Stock from normal slow-moving Push items.
4. R1-R7 Invariance: Dead Stock is an analytical view, NOT a replacement for Red List classification.
5. Dead Stock KPIs calculation (SKU count, Total Inventory Value, Aged Value / Cash at Risk).
6. Top Dead Stock by Inventory Value table ranking.
7. Excel export generating valid workbook.
"""

import io
import pytest
import pandas as pd
import openpyxl

from src.data_loader import load_data
from src.validation import validate_data
from src.business_logic import run_analysis
from src.dead_stock import (
    identify_dead_stock,
    get_dead_stock_summary,
    get_top_dead_stock_by_value,
    export_dead_stock_excel,
)

SAMPLE_PATH = "sample_data/CompanyX_Data.xlsx"


@pytest.fixture(scope="module")
def analysis_data():
    raw = load_data(SAMPLE_PATH)
    val = validate_data(*raw)
    res = run_analysis(*raw, validation=val)
    items = res["analysis_items"]
    sales = res["clean_sales"]
    dead_df = identify_dead_stock(items, sales_df=sales)
    return raw, val, res, items, sales, dead_df


def test_01_identify_dead_stock_criteria(analysis_data):
    """Verify that items in dead_stock_df have strictly 6M Sales Qty = 0."""
    _, _, _, items, _, dead_df = analysis_data

    # Every dead stock item must have 6M Sales Qty == 0
    assert (dead_df["6M Sales Qty"] == 0).all()

    # In sample data, exactly 1 active item (SKU1012) has 0 sales in the 6M period
    assert len(dead_df) == 1
    assert dead_df.iloc[0]["SKU"] == "SKU1012"
    assert dead_df.iloc[0]["Total Stock"] == 163.0
    assert dead_df.iloc[0]["Inventory Value"] == pytest.approx(55420.0)


def test_02_all_10_required_columns_present(analysis_data):
    """Verify all 10 required columns exist and are non-empty."""
    _, _, _, _, _, dead_df = analysis_data

    required_cols = [
        "SKU",
        "Description",
        "Category",
        "Total Stock",
        "Inventory Value",
        "Aged Qty",
        "Aged Value",
        "Last Available Sales Information",
        "Risk Score",
        "Recommended Action",
    ]

    for col in required_cols:
        assert col in dead_df.columns, f"Required column '{col}' missing from Dead Stock DataFrame"

    row = dead_df.iloc[0]
    assert row["SKU"] == "SKU1012"
    assert len(row["Description"]) > 0
    assert row["Category"] == "Category B"
    assert row["Total Stock"] > 0
    assert row["Inventory Value"] > 0
    assert row["Aged Qty"] == 49.0
    assert row["Aged Value"] == pytest.approx(16660.0)
    assert len(row["Last Available Sales Information"]) > 0
    assert row["Risk Score"] == 87
    assert len(row["Recommended Action"]) > 0


def test_03_separation_from_push_items(analysis_data):
    """Verify Dead Stock items are strictly separated from normal Push items."""
    _, _, _, items, _, dead_df = analysis_data

    push_items = items[items["Class"] == "Push"]
    dead_skus = set(dead_df["SKU"])
    push_skus = set(push_items["SKU"])

    # Push items have active sales and must NEVER intersect with Dead Stock (0 sales)
    assert dead_skus.isdisjoint(push_skus), "Dead Stock items must not appear in Push classification"
    assert (push_items["6M Sales Qty"] > 0).all()


def test_04_r1_to_r7_invariance(analysis_data):
    """Verify Dead Stock is an analytical view and does NOT overwrite Red List classification."""
    _, _, _, items, _, dead_df = analysis_data

    # Official R5 classification remains Red List
    assert dead_df.iloc[0]["Class"] == "Red List"

    # Official portfolio benchmark counts remain invariant
    counts = items["Class"].value_counts()
    assert counts["Red List"] == 7
    assert counts["No Action"] == 13
    assert counts["Push"] == 18


def test_05_dead_stock_kpis(analysis_data):
    """Verify calculation of Dead Stock KPIs."""
    _, _, _, items, _, dead_df = analysis_data
    tot_inv_val = float(items["Inventory Value"].sum())
    tot_aged_val = float(items["Aged Value"].sum())

    summary = get_dead_stock_summary(
        dead_df,
        total_inventory_value=tot_inv_val,
        total_aged_value=tot_aged_val,
    )

    assert summary["dead_stock_count"] == 1
    assert summary["dead_stock_inventory_value"] == pytest.approx(55420.0)
    assert summary["dead_stock_aged_value"] == pytest.approx(16660.0)
    assert summary["dead_stock_total_units"] == 163.0
    assert summary["dead_stock_aged_units"] == 49.0
    assert summary["dead_stock_pct_of_total_inv"] == pytest.approx(55420.0 / tot_inv_val * 100.0)
    assert summary["skus"] == ["SKU1012"]
    assert "1 SKU" in summary["kpi_label"]
    assert "₹55,420" in summary["kpi_label"]


def test_06_top_dead_stock_by_inventory_value(analysis_data):
    """Verify ranking of Top Dead Stock by Inventory Value descending."""
    _, _, _, _, _, dead_df = analysis_data
    top_dead = get_top_dead_stock_by_value(dead_df, top_n=5)

    assert "Rank" in top_dead.columns
    assert len(top_dead) == 1
    assert top_dead.iloc[0]["Rank"] == 1
    assert top_dead.iloc[0]["SKU"] == "SKU1012"
    assert top_dead.iloc[0]["Inventory Value"] == pytest.approx(55420.0)


def test_07_last_available_sales_information_synthetic():
    """Verify sales information lookup when historical sales exist."""
    items = pd.DataFrame([{
        "SKU": "SKU_ZERO", "Item Name": "Zero Sales Item", "Category": "Cat Z",
        "Class": "Red List", "Total Stock": 100, "Aged Qty": 50, "Unit Cost (Rs)": 100,
        "6M Sales Qty": 0, "Risk Score": 90, "Recommended Action": "Liquidate"
    }])
    sales = pd.DataFrame([{
        "Month": "Jan-2026", "CustomerID": "C100", "CustomerName": "Cust 100",
        "Industry": "Hotel", "SKU": "SKU_ZERO", "Qty": 5, "Value (Rs)": 500
    }])

    dead_df = identify_dead_stock(items, sales_df=sales)
    assert len(dead_df) == 1
    info = dead_df.iloc[0]["Last Available Sales Information"]
    assert "Jan-2026" in info
    assert "5 units" in info
    assert "₹500" in info


def test_08_export_dead_stock_excel(analysis_data):
    """Verify Excel export generates valid workbook with headers and data."""
    _, _, _, _, _, dead_df = analysis_data
    excel_bytes = export_dead_stock_excel(dead_df)
    assert len(excel_bytes) > 4000

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert "Dead Stock Analysis" in wb.sheetnames
    ws = wb["Dead Stock Analysis"]

    headers = [ws.cell(row=4, column=c).value for c in range(1, 11)]
    assert "SKU" in headers
    assert "Description" in headers
    assert "Inventory Value" in headers
    assert "Last Available Sales Information" in headers
    assert ws.max_row >= 5


def test_09_run_analysis_returns_dead_stock(analysis_data):
    """Verify that run_analysis returns dead_stock and dead_stock_summary in dict."""
    _, _, res, _, _, _ = analysis_data
    assert "dead_stock" in res
    assert "dead_stock_summary" in res
    assert isinstance(res["dead_stock"], pd.DataFrame)
    assert isinstance(res["dead_stock_summary"], dict)
    assert res["dead_stock_summary"]["dead_stock_count"] == 1


def test_10_multiple_dead_stock_items_ranking_descending():
    """Verify that when multiple items have 0 sales, they are ranked strictly by Inventory Value descending."""
    items = pd.DataFrame([
        {
            "SKU": "SKU_SMALL", "Item Name": "Small Dead Item", "Category": "Cat A",
            "Class": "Red List", "Total Stock": 10, "Aged Qty": 10, "Unit Cost (Rs)": 100,
            "Inventory Value": 1000, "Aged Value": 1000, "6M Sales Qty": 0, "Risk Score": 80,
            "Recommended Action": "Discontinue"
        },
        {
            "SKU": "SKU_LARGE", "Item Name": "Large Dead Item", "Category": "Cat B",
            "Class": "Red List", "Total Stock": 500, "Aged Qty": 200, "Unit Cost (Rs)": 200,
            "Inventory Value": 100000, "Aged Value": 40000, "6M Sales Qty": 0, "Risk Score": 95,
            "Recommended Action": "Liquidate"
        },
        {
            "SKU": "SKU_MEDIUM", "Item Name": "Medium Dead Item", "Category": "Cat C",
            "Class": "Red List", "Total Stock": 100, "Aged Qty": 50, "Unit Cost (Rs)": 150,
            "Inventory Value": 15000, "Aged Value": 7500, "6M Sales Qty": 0, "Risk Score": 85,
            "Recommended Action": "Review"
        },
    ])

    dead_df = identify_dead_stock(items)
    assert len(dead_df) == 3

    top_dead = get_top_dead_stock_by_value(dead_df, top_n=2)
    assert len(top_dead) == 2
    # First item must be SKU_LARGE (₹100,000)
    assert top_dead.iloc[0]["SKU"] == "SKU_LARGE"
    assert top_dead.iloc[0]["Inventory Value"] == 100000.0
    # Second item must be SKU_MEDIUM (₹15,000)
    assert top_dead.iloc[1]["SKU"] == "SKU_MEDIUM"
    assert top_dead.iloc[1]["Inventory Value"] == 15000.0


def test_11_dead_stock_view_integrated_in_app():
    """Verify Dead Stock Analysis view, KPIs, and table are present in app.py."""
    with open("app.py", "r", encoding="utf-8") as f:
        content = f.read()

    assert "Dead Stock Analysis View" in content
    assert "Top Dead Stock by Inventory Value" in content
    assert "Dead Stock SKUs" in content
    assert "Dead Stock Inventory Value" in content
    assert "CompanyX_Dead_Stock_Analysis.xlsx" in content
    assert "identify_dead_stock" in content

