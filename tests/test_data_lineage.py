"""
Unit tests for the Data Lineage & Audit module (src/data_lineage.py).

Verifies:
1. All 12 requested pipeline stages exist in exact sequential order:
   - INPUT
   - Validation
   - Cleaning
   - SWITCH Consolidation
   - Inactive Item Filtering
   - Inventory Metrics
   - R1-R7 Classification
   - Customer Eligibility
   - Bucket Assignment
   - Risk Score
   - Priority Score
   - Sales Action List
2. Every stage contains all required fields:
   - Number of records (records_in, records_out, breakdown)
   - Important transformations (non-empty list, plain English)
   - Warnings (list)
   - Excluded records (count and identification)
   - Reason for exclusion (governance rationale)
3. Non-technical presentation: no technical code exposed.
4. Calculations match exact benchmark numbers:
   - 744 raw records ingested
   - 6 duplicate sales lines purged
   - 3 pack-size switches consolidated
   - 1 inactive item excluded (holding 169 units)
   - 38 active catalog items evaluated
   - 18 Push, 13 No Action, 7 Red List
   - 14 eligible customers (> ₹100k spend), 16 ineligible
   - 163 customer-SKU target recommendations generated
5. Excel export generates valid workbook.
"""

import io
import pytest
import openpyxl
import pandas as pd

from src.data_loader import load_data
from src.validation import validate_data
from src.business_logic import run_analysis
from src.customer_targets import build_customer_targets
from src.data_lineage import (
    build_data_lineage,
    export_data_lineage_excel,
)

SAMPLE_PATH = "sample_data/CompanyX_Data.xlsx"

EXPECTED_STAGES = [
    "INPUT",
    "Validation",
    "Cleaning",
    "SWITCH Consolidation",
    "Inactive Item Filtering",
    "Inventory Metrics",
    "R1-R7 Classification",
    "Customer Eligibility",
    "Bucket Assignment",
    "Risk Score",
    "Priority Score",
    "Sales Action List",
]


@pytest.fixture(scope="module")
def pipeline_data():
    raw = load_data(SAMPLE_PATH)
    val = validate_data(*raw)
    analysis = run_analysis(*raw, validation=val)
    targets = build_customer_targets(
        analysis["clean_item_master"],
        analysis["clean_sales"],
        analysis["analysis_items"],
        val,
        customers=raw[3],
    )
    lineage = build_data_lineage(raw, val, analysis, targets)
    return raw, val, analysis, targets, lineage


def test_01_all_12_stages_present_in_exact_order(pipeline_data):
    """Verify all 12 stages exist and follow the exact required order."""
    _, _, _, _, lineage = pipeline_data
    stages = lineage["stages"]

    assert len(stages) == 12, f"Expected 12 stages, got {len(stages)}"
    actual_names = [s["name"] for s in stages]
    assert actual_names == EXPECTED_STAGES, f"Stage sequence mismatch:\nActual: {actual_names}\nExpected: {EXPECTED_STAGES}"


def test_02_every_stage_has_all_five_required_properties(pipeline_data):
    """
    For each stage verify:
    - Number of records
    - Important transformations
    - Warnings
    - Excluded records
    - Reason for exclusion
    """
    _, _, _, _, lineage = pipeline_data
    stages = lineage["stages"]

    for s in stages:
        stage_name = s["name"]

        # 1. Number of records
        assert "records_in" in s, f"records_in missing for {stage_name}"
        assert "records_out" in s, f"records_out missing for {stage_name}"
        assert isinstance(s["records_in"], (int, float)), f"records_in not numeric for {stage_name}"
        assert isinstance(s["records_out"], (int, float)), f"records_out not numeric for {stage_name}"
        assert s["records_in"] >= 0
        assert s["records_out"] >= 0

        # 2. Important transformations
        assert "transformations" in s, f"transformations missing for {stage_name}"
        assert isinstance(s["transformations"], list), f"transformations not list for {stage_name}"
        assert len(s["transformations"]) > 0, f"transformations empty for {stage_name}"

        # 3. Warnings
        assert "warnings" in s, f"warnings missing for {stage_name}"
        assert isinstance(s["warnings"], list), f"warnings not list for {stage_name}"

        # 4. Excluded records
        assert "excluded_records_count" in s, f"excluded_records_count missing for {stage_name}"
        assert isinstance(s["excluded_records_count"], (int, float)), f"excluded_records_count not numeric for {stage_name}"
        assert "excluded_records_detail" in s, f"excluded_records_detail missing for {stage_name}"

        # 5. Reason for exclusion
        assert "exclusion_reason" in s, f"exclusion_reason missing for {stage_name}"
        assert len(s["exclusion_reason"]) > 0, f"exclusion_reason empty for {stage_name}"


def test_03_no_technical_code_exposed_in_explanations(pipeline_data):
    """Verify that transformation descriptions and reasons do not expose python code or tracebacks."""
    _, _, _, _, lineage = pipeline_data
    stages = lineage["stages"]

    prohibited_technical_tokens = [
        "def ", "lambda ", "import ", "__", "try:", "except Exception",
        "pd.DataFrame", ".apply(", ".fillna(", "Traceback"
    ]

    for s in stages:
        combined_text = " ".join(s["transformations"] + s["warnings"] + [s["exclusion_reason"]])
        for token in prohibited_technical_tokens:
            assert token not in combined_text, f"Technical code token '{token}' exposed in stage '{s['name']}' text"


def test_04_specific_sample_data_stage_metrics(pipeline_data):
    """Verify the exact benchmark numbers for all key pipeline stages."""
    _, _, _, _, lineage = pipeline_data
    stages = {s["name"]: s for s in lineage["stages"]}

    # Stage 1: INPUT
    assert stages["INPUT"]["records_in"] == 744  # 42 + 42 + 630 + 30
    assert stages["INPUT"]["excluded_records_count"] == 0

    # Stage 2: Validation
    assert stages["Validation"]["records_in"] == 744
    assert stages["Validation"]["excluded_records_count"] == 0

    # Stage 3: Cleaning
    assert stages["Cleaning"]["records_out"] == 738  # 744 - 6 duplicates
    assert stages["Cleaning"]["excluded_records_count"] == 6
    assert "duplicate" in stages["Cleaning"]["exclusion_reason"].lower()

    # Stage 4: SWITCH Consolidation
    assert stages["SWITCH Consolidation"]["excluded_records_count"] == 3
    assert "Rule R3" in stages["SWITCH Consolidation"]["exclusion_reason"] or "R3" in stages["SWITCH Consolidation"]["exclusion_reason"]

    # Stage 5: Inactive Item Filtering
    assert stages["Inactive Item Filtering"]["excluded_records_count"] == 1
    assert "SKU1041" in stages["Inactive Item Filtering"]["excluded_records_detail"]
    assert "Rule R4" in stages["Inactive Item Filtering"]["exclusion_reason"]

    # Stage 6: Inventory Metrics
    assert stages["Inventory Metrics"]["records_out"] == 38
    assert stages["Inventory Metrics"]["excluded_records_count"] == 0

    # Stage 7: R1-R7 Classification
    assert stages["R1-R7 Classification"]["records_out"] == 38
    assert stages["R1-R7 Classification"]["record_breakdown"]["Push (Sales Intervention)"] == 18
    assert stages["R1-R7 Classification"]["record_breakdown"]["No Action (Healthy / Zero Aged)"] == 13
    assert stages["R1-R7 Classification"]["record_breakdown"]["Red List (Discontinue / Write-off)"] == 7
    assert stages["R1-R7 Classification"]["excluded_records_count"] == 20

    # Stage 8: Customer Eligibility
    assert stages["Customer Eligibility"]["records_out"] == 14
    assert stages["Customer Eligibility"]["excluded_records_count"] == 16
    assert "100,000" in stages["Customer Eligibility"]["exclusion_reason"]

    # Stage 9: Bucket Assignment
    assert stages["Bucket Assignment"]["records_out"] == 163
    assert stages["Bucket Assignment"]["record_breakdown"]["Bucket 1 (Direct Past Buyers)"] > 0
    assert stages["Bucket Assignment"]["record_breakdown"]["Bucket 2 (Industry Affinity)"] > 0
    assert stages["Bucket Assignment"]["record_breakdown"]["Bucket 3 (Category Buyers)"] > 0

    # Stage 10: Risk Score
    assert stages["Risk Score"]["records_out"] == 38
    assert stages["Risk Score"]["excluded_records_count"] == 0

    # Stage 11: Priority Score
    assert stages["Priority Score"]["records_in"] == 18
    assert stages["Priority Score"]["records_out"] == 18
    assert stages["Priority Score"]["excluded_records_count"] == 20

    # Stage 12: Sales Action List
    assert stages["Sales Action List"]["records_out"] == 163
    assert stages["Sales Action List"]["record_breakdown"]["Push Items Covered"] == 18


def test_05_matrix_df_structure(pipeline_data):
    """Verify matrix_df contains all 12 stages and expected columns."""
    _, _, _, _, lineage = pipeline_data
    matrix_df = lineage["matrix_df"]

    assert len(matrix_df) == 12
    required_cols = [
        "Stage #",
        "Pipeline Stage",
        "Records In",
        "Records Out",
        "Excluded Records",
        "Exclusion Reason",
        "Key Transformations",
        "Warnings",
    ]
    for col in required_cols:
        assert col in matrix_df.columns, f"Column '{col}' missing from matrix_df"


def test_06_excel_export(pipeline_data):
    """Verify that export_data_lineage_excel produces a valid openpyxl readable workbook."""
    _, _, _, _, lineage = pipeline_data
    excel_bytes = export_data_lineage_excel(lineage)

    assert len(excel_bytes) > 4000
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert "Data Lineage & Audit Flow" in wb.sheetnames
    ws = wb["Data Lineage & Audit Flow"]

    # Row 5 has table headers
    headers = [ws.cell(row=5, column=c).value for c in range(1, 11)]
    assert "Stage #" in headers
    assert "Pipeline Stage" in headers
    assert "Records In" in headers
    assert "Records Out" in headers
    assert "Excluded Records" in headers
    assert "Reason for Exclusion" in headers
    assert ws.max_row >= 17  # 5 header rows + 12 stage rows
