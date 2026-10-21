import io
import numpy as np
import pandas as pd
import pytest
import openpyxl

from src.business_logic import (
    generate_sku_classification_explanation,
    run_analysis,
)
from src.data_loader import load_workbook
from src.validation import validate_data
from src.export import create_excel_report
from src.customer_targets import build_customer_targets


@pytest.fixture(scope="module")
def base_results():
    """Load standard Company X sample data and run full review pipeline."""
    raw = load_workbook("sample_data/CompanyX_Data.xlsx")
    val = validate_data(*raw)
    analysis = run_analysis(*raw, validation=val)
    targets = build_customer_targets(
        analysis["clean_item_master"],
        analysis["clean_sales"],
        analysis["analysis_items"],
        val,
        customers=raw[3],
    )
    return {
        "raw": raw,
        "val": val,
        "analysis": analysis,
        "targets": targets,
        "items": analysis["analysis_items"],
    }


class TestSKUClassificationExplainability:
    """Tests for deterministic explainability panel logic across all catalog items."""

    def test_push_sku1006_exact_explanation_format(self, base_results):
        """SKU1006 must match required deterministic Push template and metrics."""
        items = base_results["items"]
        row = items[items["SKU"] == "SKU1006"].iloc[0]

        expl = generate_sku_classification_explanation(row)
        assert expl["classification"] == "Push"
        assert "SKU1006 is classified as PUSH" in expl["explanation"]
        assert "aged inventory and its months to clear exceeds the 3-month threshold" in expl["explanation"]
        assert "₹9,150 of aged inventory value" in expl["explanation"]
        assert "estimated 4.0 months to clear" in expl["explanation"]

        # Rule Triggered and Recommendation
        assert "Rule R5" in expl["rule_triggered"]
        assert "Push" in expl["rule_triggered"]
        assert len(expl["recommended_action"]) > 10

        # All 9 Supporting Metrics present
        m = expl["supporting_metrics"]
        assert m["6M Sales Qty"] == 306.0
        assert m["Master Pack"] == 10.0
        assert m["Aged Qty"] == 61.0
        assert m["Aged Value"] == 9150.0
        assert round(m["Avg Monthly Sales"], 1) == 51.0
        assert round(m["Months to Clear"], 2) == 3.98
        assert m["Inventory Value"] == 30450.0
        assert m["Risk Score"] == 41
        assert m["Risk Level"] == "Medium"

    def test_push_sku1005_explanation(self, base_results):
        """SKU1005 must explain aged inventory and clearance runway exceeding 3 months."""
        items = base_results["items"]
        row = items[items["SKU"] == "SKU1005"].iloc[0]

        expl = generate_sku_classification_explanation(row)
        assert expl["classification"] == "Push"
        assert "SKU1005 is classified as PUSH" in expl["explanation"]
        assert "₹31,280 of aged inventory value" in expl["explanation"]
        assert "estimated 6.6 months to clear" in expl["explanation"]

        m = expl["supporting_metrics"]
        assert m["6M Sales Qty"] == 280.0
        assert m["Master Pack"] == 50.0
        assert m["Aged Qty"] == 92.0
        assert m["Aged Value"] == 31280.0
        assert m["Risk Level"] == "High"
        assert m["Risk Score"] == 60

    def test_no_action_zero_aged_stock(self, base_results):
        """SKU1004 has 0 aged stock; explanation must explicitly state zero aged stock."""
        items = base_results["items"]
        row = items[items["SKU"] == "SKU1004"].iloc[0]

        expl = generate_sku_classification_explanation(row)
        assert expl["classification"] == "No Action"
        assert "NO ACTION" in expl["explanation"]
        assert "zero aged stock" in expl["explanation"].lower() or "no aged stock" in expl["explanation"].lower()
        assert "0 units" in expl["explanation"]
        assert "No aged stock" in expl["reason_summary"]
        assert "Rule R5" in expl["rule_triggered"]

        m = expl["supporting_metrics"]
        assert m["Aged Qty"] == 0.0
        assert m["Aged Value"] == 0.0

    def test_no_action_mtc_below_3(self, base_results):
        """SKU1001 has aged stock (80 units) but MTC < 3 (2.75); explanation must explicitly state MTC below 3."""
        items = base_results["items"]
        row = items[items["SKU"] == "SKU1001"].iloc[0]

        expl = generate_sku_classification_explanation(row)
        assert expl["classification"] == "No Action"
        assert "NO ACTION" in expl["explanation"]
        assert "months to clear is below 3" in expl["explanation"]
        assert "2.75" in expl["explanation"]
        assert "80 units" in expl["explanation"]
        assert expl["reason_summary"] == "Months to Clear below 3"
        assert "Rule R5" in expl["rule_triggered"]

        m = expl["supporting_metrics"]
        assert m["Aged Qty"] == 80.0
        assert m["Aged Value"] == 2000.0
        assert round(m["Months to Clear"], 2) == 2.75

    def test_no_action_sku1017_mtc_below_3(self, base_results):
        """SKU1017 has aged stock (135 units, ₹20,250) but MTC < 3 (2.57); explanation must cite MTC below 3."""
        items = base_results["items"]
        row = items[items["SKU"] == "SKU1017"].iloc[0]

        expl = generate_sku_classification_explanation(row)
        assert expl["classification"] == "No Action"
        assert "months to clear is below 3" in expl["explanation"]
        assert "2.57" in expl["explanation"]
        assert expl["reason_summary"] == "Months to Clear below 3"

    def test_red_list_zero_sales(self, base_results):
        """SKU1012 has 0 sales < 50 Master Pack; explanation must cite 6M Sales Qty < 1 Master Pack."""
        items = base_results["items"]
        row = items[items["SKU"] == "SKU1012"].iloc[0]

        expl = generate_sku_classification_explanation(row)
        assert expl["classification"] == "Red List"
        assert "RED LIST" in expl["explanation"]
        assert "6M Sales Qty < 1 Master Pack" in expl["explanation"]
        assert "0 units" in expl["explanation"]
        assert "50 units" in expl["explanation"]
        assert expl["reason_summary"] == "6M Sales Qty < 1 Master Pack"
        assert "Rule R5 — Step 1" in expl["rule_triggered"]
        assert "delisting" in expl["recommended_action"].lower() or "commercial review" in expl["recommended_action"].lower()

        m = expl["supporting_metrics"]
        assert m["6M Sales Qty"] == 0.0
        assert m["Master Pack"] == 50.0

    def test_red_list_insufficient_sales(self, base_results):
        """SKU1026 has 12 sales < 50 Master Pack; explanation must cite 6M Sales Qty < 1 Master Pack."""
        items = base_results["items"]
        row = items[items["SKU"] == "SKU1026"].iloc[0]

        expl = generate_sku_classification_explanation(row)
        assert expl["classification"] == "Red List"
        assert "6M Sales Qty < 1 Master Pack" in expl["explanation"]
        assert "12 units" in expl["explanation"]
        assert "50 units" in expl["explanation"]
        assert expl["reason_summary"] == "6M Sales Qty < 1 Master Pack"

        m = expl["supporting_metrics"]
        assert m["6M Sales Qty"] == 12.0
        assert m["Master Pack"] == 50.0

    def test_dataframe_columns_populated(self, base_results):
        """Analysis items DataFrame must have precomputed explainability columns for all rows."""
        items = base_results["items"]
        expected_cols = [
            "Classification Explanation",
            "Rule Triggered",
            "Classification Reason",
            "Recommended Action",
        ]
        for col in expected_cols:
            assert col in items.columns
            # No blank values
            assert items[col].notna().all()
            assert (items[col].str.len() > 0).all()

    def test_inactive_item_handling(self):
        """Synthetic inactive SKU test must reference Rule R4."""
        inactive_item = {
            "SKU": "SKU9999",
            "Item Name": "Discontinued Product",
            "Status": "Inactive",
            "Class": "Inactive",
            "Total Stock": 100,
            "Unit Cost (Rs)": 50,
        }
        expl = generate_sku_classification_explanation(inactive_item)
        assert "Inactive" in expl["reason_summary"]
        assert "Rule R4" in expl["rule_triggered"]

    def test_core_r1_r7_invariants_preserved(self, base_results):
        """Adding explainability panel must strictly preserve all benchmark counts and classifications."""
        items = base_results["items"]
        counts = items["Class"].value_counts()
        assert counts.get("Red List") == 7
        assert counts.get("No Action") == 13
        assert counts.get("Push") == 18

        push_aged_val = float(items[items["Class"] == "Push"]["Aged Value"].sum())
        assert push_aged_val == 163841.0

        sku_classes = items.set_index("SKU")["Class"].to_dict()
        assert sku_classes["SKU1001"] == "No Action"
        assert sku_classes["SKU1006"] == "Push"
        assert sku_classes["SKU1019"] == "Push"
        assert sku_classes["SKU1005"] == "Push"

    def test_excel_export_includes_explainability_columns(self, base_results):
        """Excel report sheets must contain the new explainability columns."""
        excel_bytes = create_excel_report(
            base_results["analysis"]["analysis_items"],
            base_results["targets"],
            base_results["val"],
            base_results["analysis"]["inactive_stock"],
            base_results["analysis"]["switch_map_df"],
        )
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        ws_items = wb["SKU Analysis"] if "SKU Analysis" in wb.sheetnames else wb["Item Analysis"]
        headers = [ws_items.cell(1, col_idx).value for col_idx in range(1, ws_items.max_column + 1)]
        assert "Rule Triggered" in headers
        assert "Classification Reason" in headers
        assert "Classification Explanation" in headers
        assert "Recommended Action" in headers

        ws_push = wb["Push Items"]
        push_headers = [ws_push.cell(1, col_idx).value for col_idx in range(1, ws_push.max_column + 1)]
        assert "Rule Triggered" in push_headers
        assert "Classification Explanation" in push_headers
