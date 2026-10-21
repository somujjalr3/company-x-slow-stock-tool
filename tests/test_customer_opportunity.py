"""
Unit Tests for Customer Opportunity Scoring and Targeting Module.

Verifies:
1. Strict adherence to Rules R6 and R7 (Buckets 1, 2, 3 and sorting invariant).
2. Customer Opportunity Score is strictly deterministic within [0, 100].
3. Scoring factors (Affinity, Spend, Category Depth, Recency).
4. Opportunity Score does NOT move customers between buckets or alter R7 priority order.
5. All required columns are present:
   Customer ID, Bucket, 6M Purchase Value, Industry, Relevant Purchase Evidence, Opportunity Score, Reason.
6. Deterministic "Why this customer?" explanation matches:
   "Bucket 2 because the customer is eligible and operates in an industry associated with previous buyers of this item."
7. Regression benchmarks on sample data (SKU1005 top 3 targets: C220 B1, C224 B1, C214 B2).
"""

import io
import pytest
import pandas as pd
import openpyxl

from src.data_loader import load_data
from src.validation import validate_data
from src.business_logic import run_analysis
from src.customer_targets import (
    build_customer_targets,
    calculate_customer_opportunity_score,
    build_customer_evidence_and_reason,
)
from src.export import create_excel_report

SAMPLE_PATH = "sample_data/CompanyX_Data.xlsx"


@pytest.fixture
def sample_analysis():
    raw = load_data(SAMPLE_PATH)
    val = validate_data(*raw)
    res = run_analysis(*raw, validation=val)
    targets = build_customer_targets(
        res["clean_item_master"],
        res["clean_sales"],
        res["analysis_items"],
        val,
        customers=raw[3],
    )
    return raw, val, res, targets


# ==============================================================================
# 1. OPPORTUNITY SCORE BOUNDS & DETERMINISM
# ==============================================================================
def test_opportunity_score_bounds_and_determinism():
    """Verify opportunity score is strictly bounded [0, 100] and deterministic."""
    # Minimum possible for eligible account
    score_min, lvl_min, bd_min = calculate_customer_opportunity_score(
        spend=100001,
        bucket=3,
        bought_exact=False,
        bought_switch=False,
        category_spend=0,
        recency_gap=5,
    )
    assert 0 <= score_min <= 100
    assert lvl_min == "Developing"
    assert bd_min["affinity_pts"] == 12
    assert bd_min["spend_pts"] == 5
    assert bd_min["cat_pts"] == 0
    assert bd_min["rec_pts"] == 3

    # Maximum possible score
    score_max, lvl_max, bd_max = calculate_customer_opportunity_score(
        spend=1000000,
        bucket=1,
        bought_exact=True,
        bought_switch=False,
        category_spend=250000,
        recency_gap=0,
    )
    assert score_max == 100
    assert lvl_max == "High"
    assert bd_max["affinity_pts"] == 35
    assert bd_max["spend_pts"] == 25
    assert bd_max["cat_pts"] == 20
    assert bd_max["rec_pts"] == 20

    # Intermediate score
    score_mid, lvl_mid, _ = calculate_customer_opportunity_score(
        spend=350000,
        bucket=2,
        bought_exact=False,
        bought_switch=False,
        category_spend=60000,
        recency_gap=1,
    )
    assert score_mid == 18 + 16 + 12 + 15  # 61
    assert lvl_mid == "Medium-High"


# ==============================================================================
# 2. DETERMINISTIC REASON GENERATOR
# ==============================================================================
def test_deterministic_reason_and_evidence():
    """Verify deterministic explanations conform strictly to specification."""
    # Bucket 2 Reason exact match to prompt specification
    ev_b2, r_b2 = build_customer_evidence_and_reason(
        bucket=2,
        category="Fasteners",
        industry="Automotive",
        spend=450000,
        bought_exact=False,
        bought_switch=False,
        exact_qty=0,
        exact_val=0,
        exact_orders=0,
        switch_qty=0,
        switch_val=0,
        switch_orders=0,
        switch_skus=[],
        cat_qty=0,
        cat_val=0,
        cat_orders=0,
        other_cat_skus=[],
        prior_buyer_names=["Customer A", "Customer B"],
        latest_active_month="Sep-2026",
    )
    assert r_b2 == "Bucket 2 because the customer is eligible and operates in an industry associated with previous buyers of this item."
    assert "Industry 'Automotive' matches previous buyers (Customer A, Customer B)" in ev_b2
    assert "last active Sep-2026" in ev_b2

    # Bucket 1 Exact item
    ev_b1, r_b1 = build_customer_evidence_and_reason(
        bucket=1,
        category="Fasteners",
        industry="Caterer",
        spend=600000,
        bought_exact=True,
        bought_switch=False,
        exact_qty=50,
        exact_val=75000,
        exact_orders=3,
        switch_qty=0,
        switch_val=0,
        switch_orders=0,
        switch_skus=[],
        cat_qty=50,
        cat_val=75000,
        cat_orders=3,
        other_cat_skus=[],
        prior_buyer_names=[],
        latest_active_month="Sep-2026",
    )
    assert r_b1 == "Bucket 1 because the customer is eligible and previously purchased this item."
    assert "Bought exact SKU: 50 units (₹75,000) across 3 order(s)" in ev_b1

    # Bucket 1 Switch partner
    ev_b1_sw, r_b1_sw = build_customer_evidence_and_reason(
        bucket=1,
        category="Fasteners",
        industry="Hotel",
        spend=300000,
        bought_exact=False,
        bought_switch=True,
        exact_qty=0,
        exact_val=0,
        exact_orders=0,
        switch_qty=40,
        switch_val=45000,
        switch_orders=2,
        switch_skus=["SKU1003"],
        cat_qty=40,
        cat_val=45000,
        cat_orders=2,
        other_cat_skus=[],
        prior_buyer_names=[],
        latest_active_month="Aug-2026",
    )
    assert r_b1_sw == "Bucket 1 because the customer is eligible and previously purchased the switch partner item."
    assert "Bought SWITCH partner (SKU1003): 40 units (₹45,000) across 2 order(s)" in ev_b1_sw

    # Bucket 3 Same Category
    ev_b3, r_b3 = build_customer_evidence_and_reason(
        bucket=3,
        category="Bearings",
        industry="Manufacturing",
        spend=220000,
        bought_exact=False,
        bought_switch=False,
        exact_qty=0,
        exact_val=0,
        exact_orders=0,
        switch_qty=0,
        switch_val=0,
        switch_orders=0,
        switch_skus=[],
        cat_qty=80,
        cat_val=60000,
        cat_orders=4,
        other_cat_skus=["SKU1010"],
        prior_buyer_names=[],
        latest_active_month="Sep-2026",
    )
    assert r_b3 == "Bucket 3 because the customer is eligible and previously purchased other items in the Bearings category."
    assert "Bought Bearings category: 80 units (₹60,000) across 4 order(s) (SKU1010)" in ev_b3


# ==============================================================================
# 3. REQUIRED COLUMNS & DATA STRUCTURE
# ==============================================================================
def test_all_required_columns_present(sample_analysis):
    """Verify that every target table contains all user-specified fields."""
    _, _, res, targets = sample_analysis
    items = res["analysis_items"]
    push_skus = items.loc[items["Class"] == "Push", "SKU"].tolist()

    required_fields = [
        "Customer ID",
        "Bucket",
        "6M Purchase Value",
        "Industry",
        "Relevant Purchase Evidence",
        "Opportunity Score",
        "Reason",
    ]

    for sku in push_skus:
        df = targets.get(sku, pd.DataFrame())
        if df.empty:
            continue
        for field in required_fields:
            assert field in df.columns, f"Field '{field}' missing from targets for {sku}"
            assert df[field].notna().all(), f"Field '{field}' has null values for {sku}"


# ==============================================================================
# 4. INVARIANT: OPPORTUNITY SCORE NEVER MOVES BUCKETS OR ALTERS R7 SORT
# ==============================================================================
def test_opportunity_score_does_not_override_r6_or_r7(sample_analysis):
    """Verify Opportunity Score never alters Bucket assignment or R7 sorting."""
    _, _, res, targets = sample_analysis
    items = res["analysis_items"]
    push_skus = items.loc[items["Class"] == "Push", "SKU"].tolist()

    for sku in push_skus:
        df = targets.get(sku, pd.DataFrame())
        if df.empty:
            continue

        # 1. Bucket sequence must strictly be non-decreasing (Bucket 1 -> 2 -> 3)
        buckets = df["Bucket"].tolist()
        assert buckets == sorted(buckets), f"Buckets out of order for {sku}"

        # 2. Within each bucket, 6M Purchase Value must strictly be descending
        for b in [1, 2, 3]:
            b_df = df[df["Bucket"] == b]
            if len(b_df) > 1:
                spends = b_df["6M Purchase Value"].tolist()
                assert spends == sorted(spends, reverse=True), f"Spend not descending in B{b} for {sku}"

        # 3. Priority must be strictly 1..N sequential
        assert df["Priority"].tolist() == list(range(1, len(df) + 1))

        # 4. Opportunity scores are valid integers between 0 and 100
        scores = df["Opportunity Score"].tolist()
        for sc in scores:
            assert isinstance(sc, (int, float))
            assert 0 <= sc <= 100


# ==============================================================================
# 5. BENCHMARK VERIFICATION: SKU1005 TARGETS
# ==============================================================================
def test_sku1005_benchmark_targets(sample_analysis):
    """
    Verify SKU1005 customer targets:
    1. Customer 220, Bucket 1 (Score 95)
    2. Customer 224, Bucket 1 (Score 87)
    3. Customer 214, Bucket 2 (Score 58)
    """
    _, _, _, targets = sample_analysis
    t_1005 = targets.get("SKU1005", pd.DataFrame())
    assert len(t_1005) == 9

    # Top target 1
    t1 = t_1005.iloc[0]
    assert t1["Priority"] == 1
    assert t1["CustomerID"] == "C220"
    assert t1["CustomerName"] == "Customer 220"
    assert t1["Bucket"] == 1
    assert t1["6M Purchase Value"] == pytest.approx(665004.0)
    assert t1["Opportunity Score"] == 95
    assert t1["Opportunity Level"] == "High"
    assert t1["Reason"] == "Bucket 1 because the customer is eligible and previously purchased this item."
    assert "Bought exact SKU: 207 units (₹373,736)" in t1["Relevant Purchase Evidence"]

    # Top target 2
    t2 = t_1005.iloc[1]
    assert t2["Priority"] == 2
    assert t2["CustomerID"] == "C224"
    assert t2["CustomerName"] == "Customer 224"
    assert t2["Bucket"] == 1
    assert t2["6M Purchase Value"] == pytest.approx(320780.0)
    assert t2["Opportunity Score"] == 87
    assert t2["Opportunity Level"] == "High"
    assert t2["Reason"] == "Bucket 1 because the customer is eligible and previously purchased this item."
    assert "Bought exact SKU: 73 units (₹97,402)" in t2["Relevant Purchase Evidence"]

    # Top target 3
    t3 = t_1005.iloc[2]
    assert t3["Priority"] == 3
    assert t3["CustomerID"] == "C214"
    assert t3["CustomerName"] == "Customer 214"
    assert t3["Bucket"] == 2
    assert t3["6M Purchase Value"] == pytest.approx(574001.0)
    assert t3["Opportunity Score"] == 58
    assert t3["Opportunity Level"] == "Moderate"
    assert t3["Reason"] == "Bucket 2 because the customer is eligible and operates in an industry associated with previous buyers of this item."
    assert "Industry 'Caterer' matches previous buyers" in t3["Relevant Purchase Evidence"]


# ==============================================================================
# 6. EXCEL EXPORT INTEGRATION
# ==============================================================================
def test_excel_export_includes_opportunity_score_and_reasons(sample_analysis):
    """Verify that Sheet 4 Customer Targets in Excel export contains Opportunity Score and Reason."""
    _, val, res, targets = sample_analysis
    excel_bytes = create_excel_report(
        items=res["analysis_items"],
        targets=targets,
        validation=val,
        inactive_stock=res["inactive_stock"],
        switches=res["switch_map_df"],
    )

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert "Customer Targets" in wb.sheetnames
    ws = wb["Customer Targets"]

    headers = [cell.value for cell in ws[1]]
    assert "Opportunity Score" in headers
    assert "Opportunity Level" in headers
    assert "Relevant Purchase Evidence" in headers
    assert "Reason" in headers

    # Read into DataFrame to verify values
    df_export = pd.read_excel(io.BytesIO(excel_bytes), sheet_name="Customer Targets")
    assert not df_export.empty
    assert (df_export["Opportunity Score"] >= 0).all()
    assert (df_export["Opportunity Score"] <= 100).all()
    assert df_export["Reason"].notna().all()
    assert df_export["Relevant Purchase Evidence"].notna().all()
