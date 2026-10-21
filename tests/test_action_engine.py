import io
import openpyxl
import pandas as pd
import pytest

from src.business_logic import (
    recommend_operational_action,
    compute_sales_trends,
    evaluate_recommended_actions,
    run_analysis,
)
from src.data_loader import load_workbook
from src.validation import validate_data
from src.export import create_excel_report
from src.customer_targets import build_customer_targets


@pytest.fixture(scope="module")
def base_data():
    """Load sample data and execute core analysis."""
    raw = load_workbook("sample_data/CompanyX_Data.xlsx")
    val = validate_data(*raw)
    analysis = run_analysis(*raw, validation=val)
    targets = build_customer_targets(
        analysis["clean_item_master"],
        analysis["clean_sales"],
        analysis["analysis_items"],
        val,
        customers=raw[3],
    )
    return {
        "raw": raw,
        "val": val,
        "analysis": analysis,
        "targets": targets,
        "items": analysis["analysis_items"],
        "sales": raw[2],
    }


class TestDeterministicRecommendedActionEngine:
    """Tests for the deterministic Recommended Action Engine."""

    def test_red_list_zero_sales_obsolete_inventory(self, base_data):
        """SKU1012 has 0 sales and aged inventory; must recommend review obsolete inventory & liquidation."""
        items = base_data["items"]
        row = items[items["SKU"] == "SKU1012"].iloc[0]

        rec = recommend_operational_action(row)
        assert rec["priority"] == "Urgent"
        assert "obsolete inventory" in rec["action"].lower()
        assert "liquidation" in rec["action"].lower() or "discontinuation" in rec["action"].lower()
        assert "zero" in rec["reason"].lower()
        assert "replenishment" in rec["reason"].lower()

    def test_red_list_sub_master_pack(self, base_data):
        """SKU1026 has 12 sales vs 50 Master Pack; must recommend review viability & stop replenishment."""
        items = base_data["items"]
        row = items[items["SKU"] == "SKU1026"].iloc[0]

        rec = recommend_operational_action(row)
        assert rec["priority"] == "High"
        assert "viability" in rec["action"].lower()
        assert "stop unnecessary replenishment" in rec["action"].lower()
        assert "master pack" in rec["reason"].lower() or "12" in rec["reason"]

    def test_no_action_zero_aged_stock(self, base_data):
        """SKU1004 has 0 aged stock and rapid clearance runway; no sales intervention required."""
        items = base_data["items"]
        row = items[items["SKU"] == "SKU1004"].iloc[0]

        rec = recommend_operational_action(row)
        assert rec["priority"] == "Low"
        assert "no immediate sales intervention" in rec["action"].lower()
        assert "continue monitoring" in rec["action"].lower()
        assert "zero aged stock" in rec["reason"].lower()

    def test_no_action_minor_aged_stock(self, base_data):
        """SKU1001 has 80 aged units but MTC < 3; continue monitoring inventory movement."""
        items = base_data["items"]
        row = items[items["SKU"] == "SKU1001"].iloc[0]

        rec = recommend_operational_action(row)
        assert rec["priority"] == "Low"
        assert "continue monitoring inventory movement" in rec["action"].lower()
        assert "organically" in rec["reason"].lower() or "2.75" in rec["reason"]

    def test_push_compound_critical_sku1005(self, base_data):
        """SKU1005 has high aged value, high MTC, high inventory value, and declining sales -> Urgent priority."""
        items = base_data["items"]
        row = items[items["SKU"] == "SKU1005"].iloc[0]

        rec = recommend_operational_action(row, is_declining=True, sales_trend_pct=-22.2)
        assert rec["priority"] == "Urgent"
        assert "prioritize high-value customers" in rec["action"].lower()
        assert "push through targeted sales activity" in rec["action"].lower()
        assert "high aged value" in rec["reason"].lower()
        assert "high months to clear" in rec["reason"].lower() or "runway" in rec["reason"].lower()
        assert "declining" in rec["reason"].lower()

    def test_push_high_inventory_value_sku1019(self, base_data):
        """SKU1019 has high inventory value (₹54,420); prioritize high-value customers."""
        items = base_data["items"]
        row = items[items["SKU"] == "SKU1019"].iloc[0]

        rec = recommend_operational_action(row, is_declining=False)
        assert rec["priority"] == "High"
        assert "create sales outreach list" in rec["action"].lower()
        assert "prioritize high-value customers" in rec["action"].lower()
        assert "inventory value" in rec["reason"].lower()

    def test_push_high_mtc_low_velocity_sku1038(self, base_data):
        """SKU1038 has high clearance runway (7.5m) and low sales velocity (27.2 u/m); targeted sales push."""
        items = base_data["items"]
        row = items[items["SKU"] == "SKU1038"].iloc[0]

        rec = recommend_operational_action(row, is_declining=True)
        assert rec["priority"] == "High"
        assert "push through targeted sales activity" in rec["action"].lower()
        assert "clearance runway" in rec["reason"].lower() or "velocity" in rec["reason"].lower()

    def test_push_moderate_runway_sku1006(self, base_data):
        """SKU1006 has moderate aged value (₹9,150) and 4.0 MTC; Medium priority with progress monitoring."""
        items = base_data["items"]
        row = items[items["SKU"] == "SKU1006"].iloc[0]

        rec = recommend_operational_action(row, is_declining=False)
        assert rec["priority"] == "Medium"
        assert "push through targeted sales activity" in rec["action"].lower()
        assert "monitor clearance progress" in rec["action"].lower()
        assert "moderate" in rec["reason"].lower()

    def test_sales_trends_computation(self, base_data):
        """compute_sales_trends must detect declining demand for key SKUs."""
        sales = base_data["sales"]
        switches = base_data["analysis"]["switch_map"]
        trends = compute_sales_trends(sales, switch_map=switches)

        assert isinstance(trends, dict)
        assert len(trends) > 0
        assert "SKU1005" in trends
        # SKU1005 early sales vs late sales
        assert trends["SKU1005"]["declining"] is True

    def test_evaluate_recommended_actions_dataframe(self, base_data):
        """evaluate_recommended_actions must enrich DataFrame with required columns."""
        items = base_data["items"]
        sales = base_data["sales"]
        switches = base_data["analysis"]["switch_map"]

        enriched = evaluate_recommended_actions(items, sales_lines=sales, switch_map=switches)
        assert "Recommended Action" in enriched.columns
        assert "Priority" in enriched.columns
        assert "Reason" in enriched.columns

        assert enriched["Recommended Action"].notna().all()
        assert enriched["Priority"].isin(["Urgent", "High", "Medium", "Low"]).all()
        assert enriched["Reason"].str.len().gt(10).all()

    def test_r1_r7_classification_invariants_preserved(self, base_data):
        """Action engine must never modify R1-R7 classification counts or aged values."""
        items = base_data["items"]
        counts = items["Class"].value_counts()
        assert counts.get("Red List") == 7
        assert counts.get("No Action") == 13
        assert counts.get("Push") == 18

        push_val = float(items[items["Class"] == "Push"]["Aged Value"].sum())
        assert push_val == 163841.0

    def test_export_includes_action_and_priority(self, base_data):
        """Excel export must include Priority and Recommended Action columns."""
        excel_bytes = create_excel_report(
            base_data["analysis"]["analysis_items"],
            base_data["targets"],
            base_data["val"],
            base_data["analysis"]["inactive_stock"],
            base_data["analysis"]["switch_map_df"],
        )
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        ws = wb["SKU Analysis"] if "SKU Analysis" in wb.sheetnames else wb["Item Analysis"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        assert "Priority" in headers
        assert "Recommended Action" in headers
        assert "Reason" in headers
