"""
Comprehensive 28-Point End-to-End QA Test Suite for Company X Slow Stock Action Tool.

Systematically audits and verifies all 28 required QA test points:
 1. Combined workbook upload
 2. Four-file upload
 3. Missing sheet handling
 4. Missing column handling
 5. Invalid data handling
 6. Duplicate data handling
 7. SWITCH item consolidation
 8. Inactive item isolation
 9. Red List classification
10. No Action classification
11. Push classification
12. Customer Bucket 1
13. Customer Bucket 2
14. Customer Bucket 3
15. Customer appearing in multiple potential buckets (Mutual Exclusivity)
16. Missing Industry exclusion from Bucket 2
17. Risk Score deterministic calculation
18. Working Capital calculations
19. Sales Trend classification
20. Priority Score ranking
21. Customer Opportunity Score
22. What-if simulator non-destructiveness & reversibility
23. Dead Stock zero-demand isolation
24. Excel export 12-sheet integrity
25. Sales Action Center dispatch ledger & filters
26. Sales Tracker execution & logging
27. Dashboard overview KPIs & exact case regression benchmarks
28. Data Quality audit page & 12-step lineage

Official Case Regression Benchmarks Verified:
- Red List = 7
- No Action = 13
- Push = 18
- Push Aged Value = ₹163,841
- SKU1001 = No Action
- SKU1006 = Push
- SKU1019 = Push
- SKU1005 = Push
- SKU1005 top targets:
  1. Customer 220, Bucket 1
  2. Customer 224, Bucket 1
  3. Customer 214, Bucket 2
"""

import io
import pytest
import numpy as np
import pandas as pd
import openpyxl

from src.data_loader import load_workbook, load_separate_files, load_data
from src.validation import validate_data
from src.business_logic import run_analysis, build_audit_trail, calculate_working_capital_impact
from src.customer_targets import build_customer_targets
from src.sales_priority import calculate_sales_priority, get_ranked_push_table
from src.sales_trends import calculate_monthly_sales_trends
from src.scenario_simulator import simulate_sku_sales_uplift, simulate_portfolio_uplift
from src.dead_stock import identify_dead_stock, get_dead_stock_summary, get_top_dead_stock_by_value
from src.data_lineage import build_data_lineage
from src.sales_action_center import (
    build_sales_action_table,
    filter_sales_action_table,
    sort_sales_action_table,
    generate_customer_sku_action,
    export_sales_action_excel,
    init_sales_tracker,
    update_sales_tracker_lead,
    calculate_sales_tracker_kpis,
    export_sales_tracker_excel,
    TRACKER_STATUS_OPTIONS,
)
from src.export import create_excel_report
from src.config import get_default_config

SAMPLE_PATH = "sample_data/CompanyX_Data.xlsx"


@pytest.fixture(scope="module")
def benchmark_data():
    raw = load_workbook(SAMPLE_PATH)
    val = validate_data(*raw)
    res = run_analysis(*raw, validation=val)
    targets = build_customer_targets(
        res["clean_item_master"],
        res["clean_sales"],
        res["analysis_items"],
        val,
        customers=raw[3],
    )
    items = res["analysis_items"]
    action_table = build_sales_action_table(items, targets, sales_trends=res.get("sales_trends"))
    return {
        "raw": raw,
        "val": val,
        "res": res,
        "targets": targets,
        "items": items,
        "action_table": action_table,
    }


# ==============================================================================
# TEST 1: COMBINED WORKBOOK UPLOAD
# ==============================================================================
def test_01_combined_workbook_upload():
    """Verify loading single Excel file containing all 4 required sheets."""
    raw = load_workbook(SAMPLE_PATH)
    im, stock, sales, cust = raw
    assert len(im) == 42, f"Expected 42 Item Master rows, got {len(im)}"
    assert len(stock) == 42, f"Expected 42 Stock Aging rows, got {len(stock)}"
    assert len(sales) == 630, f"Expected 630 Sales rows, got {len(sales)}"
    assert len(cust) == 30, f"Expected 30 Customer rows, got {len(cust)}"
    assert "SKU" in im.columns
    assert "Qty 91-180 days" in stock.columns
    assert "CustomerID" in sales.columns
    assert "CustomerName" in cust.columns


# ==============================================================================
# TEST 2: FOUR-FILE UPLOAD
# ==============================================================================
def test_02_four_file_upload():
    """Verify loading 4 separate Excel file streams yields identical schema."""
    raw = load_workbook(SAMPLE_PATH)
    im_raw, stock_raw, sales_raw, cust_raw = raw

    # Convert to in-memory Excel files
    im_buf = io.BytesIO()
    with pd.ExcelWriter(im_buf, engine="openpyxl") as writer:
        im_raw.to_excel(writer, sheet_name="Item_Master", index=False)
    im_buf.seek(0)
    im_buf.name = "Item_Master.xlsx"

    stock_buf = io.BytesIO()
    with pd.ExcelWriter(stock_buf, engine="openpyxl") as writer:
        stock_raw.to_excel(writer, sheet_name="Stock_Aging", index=False)
    stock_buf.seek(0)
    stock_buf.name = "Stock_Aging.xlsx"

    sales_buf = io.BytesIO()
    with pd.ExcelWriter(sales_buf, engine="openpyxl") as writer:
        sales_raw.to_excel(writer, sheet_name="Sales_Lines", index=False)
    sales_buf.seek(0)
    sales_buf.name = "Sales_Lines.xlsx"

    cust_buf = io.BytesIO()
    with pd.ExcelWriter(cust_buf, engine="openpyxl") as writer:
        cust_raw.to_excel(writer, sheet_name="Customers", index=False)
    cust_buf.seek(0)
    cust_buf.name = "Customers.xlsx"

    loaded_im, loaded_stock, loaded_sales, loaded_cust = load_separate_files(
        im_buf, stock_buf, sales_buf, cust_buf
    )
    assert len(loaded_im) == 42
    assert len(loaded_stock) == 42
    assert len(loaded_sales) == 630
    assert len(loaded_cust) == 30


# ==============================================================================
# TEST 3: MISSING SHEET HANDLING
# ==============================================================================
def test_03_missing_sheet_handling():
    """Verify loading workbook with a missing required sheet raises informative error."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Item_Master"
    ws.append(["SKU", "Item Name", "Category", "Master Pack Qty", "Unit Cost (Rs)", "Status"])
    ws.append(["SKU999", "Widget", "Category A", 10, 50, "Active"])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    with pytest.raises(ValueError) as exc_info:
        load_workbook(buf)
    err_msg = str(exc_info.value).lower()
    assert "missing required sheet" in err_msg


# ==============================================================================
# TEST 4: MISSING COLUMN HANDLING
# ==============================================================================
def test_04_missing_column_handling():
    """Verify validation detects missing mandatory columns and raises informative error."""
    raw = load_workbook(SAMPLE_PATH)
    im, stock, sales, cust = [df.copy() for df in raw]
    im_bad = im.drop(columns=["Master Pack Qty"])

    with pytest.raises(ValueError, match="Item_Master is missing columns"):
        validate_data(im_bad, stock, sales, cust)


# ==============================================================================
# TEST 5: INVALID DATA HANDLING
# ==============================================================================
def test_05_invalid_data_handling():
    """Verify pipeline gracefully handles negative quantities and non-numeric costs."""
    raw = load_workbook(SAMPLE_PATH)
    im, stock, sales, cust = [df.copy() for df in raw]
    im.loc[0, "Master Pack Qty"] = -5
    stock.loc[0, "Qty 0-30 days"] = -10
    sales.loc[0, "Qty"] = -1

    val = validate_data(im, stock, sales, cust)
    assert len(val["issues"]) > 0
    problems = [i["Problem"] for i in val["issues"]]
    assert any("Invalid numeric values" in p for p in problems)

    # run_analysis must execute cleanly without unhandled exception
    res = run_analysis(im, stock, sales, cust, validation=val)
    assert not res["analysis_items"].empty


# ==============================================================================
# TEST 6: DUPLICATE DATA HANDLING
# ==============================================================================
def test_06_duplicate_data_handling(benchmark_data):
    """Verify duplicate sales rows are detected and deduplicated."""
    val = benchmark_data["val"]
    assert val["duplicate_sales_rows"] == 6, f"Expected 6 duplicate sales rows, got {val['duplicate_sales_rows']}"

    clean_sales = benchmark_data["res"]["clean_sales"]
    # Original sales had 630 rows, after dropping 6 duplicates: 624 rows
    assert len(clean_sales) == 624, f"Expected 624 deduplicated sales rows, got {len(clean_sales)}"


# ==============================================================================
# TEST 7: SWITCH ITEM CONSOLIDATION (RULE R3)
# ==============================================================================
def test_07_switch_consolidation(benchmark_data):
    """Verify 3 SWITCH mappings roll up sales to new active SKU and exclude old SKUs."""
    res = benchmark_data["res"]
    switches = res["switch_map_df"]
    assert len(switches) == 3, f"Expected 3 SWITCH mappings, got {len(switches)}"

    switch_pairs = set(zip(switches["Old SKU"], switches["New SKU"]))
    assert ("SKU1003", "SKU1006") in switch_pairs
    assert ("SKU1015", "SKU1019") in switch_pairs
    assert ("SKU1030", "SKU1033") in switch_pairs

    items = res["analysis_items"]
    active_skus = set(items["SKU"].tolist())
    assert "SKU1003" not in active_skus
    assert "SKU1015" not in active_skus
    assert "SKU1030" not in active_skus
    assert "SKU1006" in active_skus
    assert "SKU1019" in active_skus
    assert "SKU1033" in active_skus


# ==============================================================================
# TEST 8: INACTIVE ITEM ISOLATION (RULE R4)
# ==============================================================================
def test_08_inactive_item_isolation(benchmark_data):
    """Verify SKU1041 is inactive with 169 units valued at ₹14,365 and isolated."""
    res = benchmark_data["res"]
    inactive = res["inactive_stock"]
    assert len(inactive) == 1, f"Expected 1 inactive item, got {len(inactive)}"
    assert inactive.iloc[0]["SKU"] == "SKU1041"
    assert inactive.iloc[0]["Total Stock"] == 169
    assert inactive.iloc[0]["Total Stock Value"] == pytest.approx(14365.0, abs=1.0)

    # Excluded from active review items
    assert "SKU1041" not in res["analysis_items"]["SKU"].values


# ==============================================================================
# TEST 9: RED LIST CLASSIFICATION (RULE R1)
# ==============================================================================
def test_09_red_list_classification(benchmark_data):
    """Verify Red List items have 6M Sales < 1 Master Pack Qty and count == 7."""
    items = benchmark_data["items"]
    red_items = items[items["Class"] == "Red List"]
    assert len(red_items) == 7, f"Expected 7 Red List items, got {len(red_items)}"

    for _, row in red_items.iterrows():
        assert row["6M Sales Qty"] < row["Master Pack Qty"], (
            f"Red List item {row['SKU']} has 6M sales {row['6M Sales Qty']} >= Master Pack {row['Master Pack Qty']}"
        )


# ==============================================================================
# TEST 10: NO ACTION CLASSIFICATION (RULE R2)
# ==============================================================================
def test_10_no_action_classification(benchmark_data):
    """Verify No Action items have Months to Clear < 3.0 and count == 13."""
    items = benchmark_data["items"]
    noaction_items = items[items["Class"] == "No Action"]
    assert len(noaction_items) == 13, f"Expected 13 No Action items, got {len(noaction_items)}"

    for _, row in noaction_items.iterrows():
        assert row["Months to Clear"] < 3.0, (
            f"No Action item {row['SKU']} has Months to Clear {row['Months to Clear']} >= 3.0"
        )


# ==============================================================================
# TEST 11: PUSH CLASSIFICATION
# ==============================================================================
def test_11_push_classification(benchmark_data):
    """Verify Push items have Aged Qty > 0, Months to Clear >= 3.0, and count == 18."""
    items = benchmark_data["items"]
    push_items = items[items["Class"] == "Push"]
    assert len(push_items) == 18, f"Expected 18 Push items, got {len(push_items)}"

    for _, row in push_items.iterrows():
        assert row["Aged Qty"] > 0, f"Push item {row['SKU']} has 0 aged quantity"
        assert row["Months to Clear"] >= 3.0, f"Push item {row['SKU']} has Months to Clear < 3.0"
        assert row["6M Sales Qty"] >= row["Master Pack Qty"], f"Push item {row['SKU']} failed Rule R1"


# ==============================================================================
# TEST 12: CUSTOMER BUCKET 1 (RULE R6 & R7)
# ==============================================================================
def test_12_customer_bucket_1(benchmark_data):
    """Verify Bucket 1 contains repeat buyers with 6M spend > ₹100,000."""
    targets = benchmark_data["targets"]
    assert "SKU1005" in targets
    t_1005 = targets["SKU1005"]
    b1_1005 = t_1005[t_1005["Bucket"] == 1]
    assert len(b1_1005) >= 2
    for _, row in b1_1005.iterrows():
        assert row["6M Purchase Value"] > 100000.0
        assert row["Bucket"] == 1


# ==============================================================================
# TEST 13: CUSTOMER BUCKET 2 (RULE R6 & R7)
# ==============================================================================
def test_13_customer_bucket_2(benchmark_data):
    """Verify Bucket 2 contains similar business accounts with spend > ₹100k."""
    targets = benchmark_data["targets"]
    t_1005 = targets["SKU1005"]
    b2_1005 = t_1005[t_1005["Bucket"] == 2]
    assert len(b2_1005) > 0
    for _, row in b2_1005.iterrows():
        assert row["6M Purchase Value"] > 100000.0
        assert row["Bucket"] == 2
        assert pd.notna(row["Industry"]) and str(row["Industry"]).strip() != ""


# ==============================================================================
# TEST 14: CUSTOMER BUCKET 3 (RULE R6 & R7)
# ==============================================================================
def test_14_customer_bucket_3(benchmark_data):
    """Verify Bucket 3 contains category cross-sell accounts with spend > ₹100k."""
    targets = benchmark_data["targets"]
    t_1005 = targets["SKU1005"]
    b3_1005 = t_1005[t_1005["Bucket"] == 3]
    assert len(b3_1005) > 0
    for _, row in b3_1005.iterrows():
        assert row["6M Purchase Value"] > 100000.0
        assert row["Bucket"] == 3


# ==============================================================================
# TEST 15: MUTUAL EXCLUSIVITY OF BUCKETS
# ==============================================================================
def test_15_customer_multiple_potential_buckets_mutual_exclusivity(benchmark_data):
    """Verify accounts never duplicate across buckets for the same SKU."""
    targets = benchmark_data["targets"]
    for sku, t_df in targets.items():
        if t_df is not None and not t_df.empty:
            assert t_df["CustomerID"].duplicated().sum() == 0, (
                f"Duplicate Customer IDs found in targeting list for {sku}"
            )


# ==============================================================================
# TEST 16: MISSING INDUSTRY EXCLUSION FROM BUCKET 2
# ==============================================================================
def test_16_missing_industry_exclusion_from_bucket_2():
    """Verify accounts with missing or blank Industry are excluded from Bucket 2."""
    im = pd.DataFrame([{
        "SKU": "SKU_P", "Item Name": "Product P", "Category": "Cat A",
        "Master Pack Qty": 10, "Unit Cost (Rs)": 50, "Alternate Item (name)": None,
        "Alternate Flag": None, "Status": "Active"
    }])
    stock = pd.DataFrame([{
        "SKU": "SKU_P",
        "Qty 0-30 days": 10, "Qty 31-60 days": 0, "Qty 61-90 days": 0,
        "Qty 91-180 days": 25, "Qty 180+ days": 25
    }])
    sales = pd.DataFrame([
        {"Month": "Apr-2026", "CustomerID": "C1", "CustomerName": "Cust 1", "Industry": "Cafe", "SKU": "SKU_P", "Qty": 20, "Value (Rs)": 150000},
        {"Month": "Apr-2026", "CustomerID": "C2", "CustomerName": "Cust 2", "Industry": None, "SKU": "OTHER_SKU", "Qty": 20, "Value (Rs)": 150000},
    ])
    cust = pd.DataFrame([
        {"CustomerID": "C1", "CustomerName": "Cust 1", "Industry": "Cafe", "Area": "North"},
        {"CustomerID": "C2", "CustomerName": "Cust 2", "Industry": "", "Area": "North"},
    ])

    res = run_analysis(im, stock, sales, cust)
    targets = build_customer_targets(res["clean_item_master"], res["clean_sales"], res["analysis_items"], customers=cust)

    t_df = targets["SKU_P"]
    c2_rows = t_df[t_df["CustomerID"] == "C2"]
    if not c2_rows.empty:
        assert c2_rows.iloc[0]["Bucket"] != 2, "Customer with missing industry was improperly assigned to Bucket 2"


# ==============================================================================
# TEST 17: RISK SCORE DETERMINISTIC SCORING
# ==============================================================================
def test_17_risk_score(benchmark_data):
    """Verify Risk Score is bounded in [0, 100] and matches risk tiers."""
    items = benchmark_data["items"]
    assert "Risk Score" in items.columns
    assert "Risk Level" in items.columns

    for _, row in items.iterrows():
        score = row["Risk Score"]
        assert 0 <= score <= 100, f"Risk Score {score} out of bounds for {row['SKU']}"
        lvl = row["Risk Level"]
        assert lvl in ["Critical", "High", "Medium", "Low"]


# ==============================================================================
# TEST 18: WORKING CAPITAL CALCULATIONS
# ==============================================================================
def test_18_working_capital_calculations(benchmark_data):
    """Verify Push Aged Value is exactly ₹163,841 and working capital structure."""
    items = benchmark_data["items"]
    push_aged = float(items.loc[items["Class"] == "Push", "Aged Value"].sum())
    assert push_aged == pytest.approx(163841.0, abs=1.0), f"Expected ₹163,841, got {push_aged}"

    wc = calculate_working_capital_impact(items)
    summary = wc["summary"]
    assert summary["push_aged_value"] == pytest.approx(163841.0, abs=1.0)
    assert summary["total_inventory_value"] > summary["push_aged_value"]


# ==============================================================================
# TEST 19: SALES TREND CLASSIFICATION
# ==============================================================================
def test_19_sales_trend(benchmark_data):
    """Verify monthly sales trends are calculated and categorized."""
    clean_sales = benchmark_data["res"]["clean_sales"]
    trends_res = calculate_monthly_sales_trends(clean_sales)
    assert trends_res["period_count"] == 6

    trends_by_sku = trends_res["trends_by_sku"]
    assert len(trends_by_sku) > 0

    valid_trends = {"Growing", "Declining", "Stable", "No Sales", "Insufficient Data"}
    for sku, t_data in trends_by_sku.items():
        assert t_data["sales_trend"] in valid_trends
        assert isinstance(t_data["monthly_qty"], dict)


# ==============================================================================
# TEST 20: PRIORITY SCORE RANKING
# ==============================================================================
def test_20_priority_score_ranking(benchmark_data):
    """Verify Priority Score ranks Push items deterministically."""
    items = benchmark_data["items"]
    ranked = get_ranked_push_table(items)
    assert len(ranked) == 18

    # Ensure strictly sorted descending by Priority Score
    scores = ranked["Priority Score"].tolist()
    assert scores == sorted(scores, reverse=True)


# ==============================================================================
# TEST 21: CUSTOMER OPPORTUNITY SCORE
# ==============================================================================
def test_21_customer_opportunity_score(benchmark_data):
    """Verify Customer Opportunity Score bounds [0, 100] and ranks targets."""
    targets = benchmark_data["targets"]
    for sku, t_df in targets.items():
        if t_df is not None and not t_df.empty:
            assert "Opportunity Score" in t_df.columns
            for _, row in t_df.iterrows():
                opp = row["Opportunity Score"]
                assert 0 <= opp <= 100


# ==============================================================================
# TEST 22: WHAT-IF SCENARIO SIMULATOR (REVERSIBLE & NON-DESTRUCTIVE)
# ==============================================================================
def test_22_what_if_simulator_reversibility(benchmark_data):
    """Verify scenario simulator calculates uplifts without modifying actual data."""
    items = benchmark_data["items"].copy()
    initial_push_count = int((items["Class"] == "Push").sum())

    sim_res = simulate_sku_sales_uplift(items.iloc[0], uplift_pct=25.0)
    assert sim_res["simulated_avg_monthly_sales"] > sim_res["current_avg_monthly_sales"]
    assert sim_res["simulated_months_to_clear"] <= sim_res["current_months_to_clear"]

    # Verify original items DataFrame is untouched
    assert int((items["Class"] == "Push").sum()) == initial_push_count
    assert items.iloc[0]["Avg Monthly Sales"] == pytest.approx(sim_res["current_avg_monthly_sales"])


# ==============================================================================
# TEST 23: DEAD STOCK ZERO-DEMAND ISOLATION
# ==============================================================================
def test_23_dead_stock_zero_demand(benchmark_data):
    """Verify Dead Stock isolates 6M Sales == 0 without affecting R1-R7."""
    items = benchmark_data["items"]
    dead_df = identify_dead_stock(items, sales_df=benchmark_data["res"]["clean_sales"])
    assert (dead_df["6M Sales Qty"] == 0).all()

    summary = get_dead_stock_summary(dead_df)
    assert summary["dead_stock_count"] == len(dead_df)
    assert summary["dead_stock_count"] == 1
    assert "SKU1012" in summary["skus"]


# ==============================================================================
# TEST 24: EXCEL EXPORT 12-SHEET INTEGRITY
# ==============================================================================
def test_24_excel_export_12_sheets(benchmark_data):
    """Verify professional Excel workbook contains all 12 required sheets."""
    res = benchmark_data["res"]
    targets = benchmark_data["targets"]
    val = benchmark_data["val"]

    excel_bytes = create_excel_report(
        items=res["analysis_items"],
        targets=targets,
        validation=val,
        inactive_stock=res["inactive_stock"],
        switches=res["switch_map_df"],
    )
    assert len(excel_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    expected_sheets = [
        "Executive Summary",
        "SKU Analysis",
        "Push Items",
        "Red List",
        "No Action",
        "Customer Targets",
        "Sales Action List",
        "Risk Analysis",
        "Working Capital",
        "Dead Stock",
        "Data Quality",
        "Assumptions & Rules",
    ]
    assert wb.sheetnames == expected_sheets


# ==============================================================================
# TEST 25: SALES ACTION CENTER DISPATCH QUEUE & FILTERS
# ==============================================================================
def test_25_sales_action_center_dispatch_queue(benchmark_data):
    """Verify consolidated cross-portfolio sales action ledger and filters."""
    table = benchmark_data["action_table"]
    assert not table.empty

    req_cols = [
        "Priority", "SKU", "Item Description", "Customer ID", "Customer Industry",
        "Bucket", "6M Purchase Value", "Opportunity Score", "Aged Value",
        "Months to Clear", "Recommended Action", "Reason"
    ]
    for c in req_cols:
        assert c in table.columns, f"Missing required Sales Action column: {c}"

    # Filter test: Bucket 1
    b1_filtered = filter_sales_action_table(table, customer_bucket="Bucket 1 (Bought Before)")
    assert (b1_filtered["Bucket Number"] == 1).all()

    # Sort test
    sorted_df = sort_sales_action_table(table, sort_by="Priority Score (High to Low)")
    assert sorted_df.iloc[0]["Priority Score"] >= sorted_df.iloc[-1]["Priority Score"]


# ==============================================================================
# TEST 26: SALES TRACKER EXECUTION & LOGGING
# ==============================================================================
def test_26_sales_tracker_execution(benchmark_data):
    """Verify Sales Outreach Tracker initializes leads, updates status, and computes KPIs."""
    table = benchmark_data["action_table"]
    tracker = init_sales_tracker(table)
    assert len(tracker) == len(table)
    assert "Call Status" in tracker.columns
    assert (tracker["Call Status"] == "Not Contacted").all()

    # Log outreach for first lead
    first_lead = tracker.iloc[0]
    sku = first_lead["SKU"]
    cid = first_lead["Customer ID"]

    updated = update_sales_tracker_lead(
        tracker,
        sku=sku,
        customer_id=cid,
        status="Order Placed (Liquidated)",
        notes="Customer accepted 25 units for immediate delivery.",
        units=25,
        value=50000.0,
    )
    kpis = calculate_sales_tracker_kpis(updated)
    assert kpis["orders_placed_count"] == 1
    assert kpis["liquidated_units"] == 25
    assert kpis["liquidated_value"] == 50000.0
    assert kpis["conversion_rate_pct"] == 100.0

    # Excel export test
    excel_bytes = export_sales_tracker_excel(updated)
    assert len(excel_bytes) > 0
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert "Sales Execution Tracker" in wb.sheetnames


# ==============================================================================
# TEST 27: DASHBOARD OVERVIEW KPIS & BENCHMARKS
# ==============================================================================
def test_27_dashboard_overview_benchmarks(benchmark_data):
    """Verify exact case regression benchmarks match required counts and targets."""
    items = benchmark_data["items"]

    # Benchmark counts
    counts = items["Class"].value_counts()
    red_count = int(counts.get("Red List", 0))
    noaction_count = int(counts.get("No Action", 0))
    push_count = int(counts.get("Push", 0))
    push_aged_value = float(items.loc[items["Class"] == "Push", "Aged Value"].sum())

    assert red_count == 7, f"Red List count must be 7, got {red_count}"
    assert noaction_count == 13, f"No Action count must be 13, got {noaction_count}"
    assert push_count == 18, f"Push count must be 18, got {push_count}"
    assert push_aged_value == pytest.approx(163841.0, abs=1.0), (
        f"Push Aged Value must be ₹163,841, got {push_aged_value}"
    )

    # Specific SKU Classifications
    sku_classes = items.set_index("SKU")["Class"].to_dict()
    assert sku_classes.get("SKU1001") == "No Action"
    assert sku_classes.get("SKU1006") == "Push"
    assert sku_classes.get("SKU1019") == "Push"
    assert sku_classes.get("SKU1005") == "Push"

    # SKU1005 Top 3 Customer Targets
    targets = benchmark_data["targets"]
    assert "SKU1005" in targets
    t_1005 = targets["SKU1005"]
    assert len(t_1005) >= 3

    top3_custs = t_1005.head(3)["CustomerName"].tolist()
    top3_buckets = t_1005.head(3)["Bucket"].tolist()

    assert top3_custs[0] == "Customer 220", f"Expected Customer 220, got {top3_custs[0]}"
    assert top3_buckets[0] == 1, f"Expected Bucket 1, got {top3_buckets[0]}"

    assert top3_custs[1] == "Customer 224", f"Expected Customer 224, got {top3_custs[1]}"
    assert top3_buckets[1] == 1, f"Expected Bucket 1, got {top3_buckets[1]}"

    assert top3_custs[2] == "Customer 214", f"Expected Customer 214, got {top3_custs[2]}"
    assert top3_buckets[2] == 2, f"Expected Bucket 2, got {top3_buckets[2]}"


# ==============================================================================
# TEST 28: DATA QUALITY AUDIT PAGE & 12-STEP LINEAGE
# ==============================================================================
def test_28_data_quality_audit_lineage(benchmark_data):
    """Verify Data Quality audit trail, 12-step lineage, and integrity metrics."""
    raw = benchmark_data["raw"]
    val = benchmark_data["val"]
    res = benchmark_data["res"]
    targets = benchmark_data["targets"]

    audit = build_audit_trail(raw, val, res, targets)
    assert len(audit["table"]) == 11
    assert audit["metrics"]["uploaded_records"] == 42 + 42 + 630 + 30
    assert audit["metrics"]["duplicates_removed"] == 6
    assert audit["metrics"]["switch_mappings"] == 3
    assert audit["metrics"]["inactive_items_excluded"] == 1
    assert audit["metrics"]["push_items"] == 18

    # Lineage verification
    lineage = build_data_lineage(raw, val, res, targets)
    assert len(lineage["stages"]) == 12
    assert lineage["summary"]["total_sales_recommendations"] > 0
